from datetime import timedelta
from unittest.mock import patch

from django.contrib.auth import get_user_model
from django.test import Client, TestCase
from django.urls import reverse
from django.utils import timezone

from inventario.models import Activo, Asignacion
from usuarios.models import Usuario

CustomUser = get_user_model()


class _FakeHTML:
    def __init__(self, *args, **kwargs):
        pass

    def write_pdf(self, *args, **kwargs):
        return b"%PDF-1.4 fake"


class _FakeWeasyprint:
    HTML = _FakeHTML


def _fake_weasyprint():
    return _FakeWeasyprint()


def _user(rol):
    return CustomUser.objects.create_user(
        username=f"rpt_{rol}_{CustomUser.objects.count()}",
        password="x",
        rol=rol,
    )


class ReportsLoginTests(TestCase):
    def test_index_requiere_login(self):
        resp = self.client.get(reverse("reports:index"))
        self.assertEqual(resp.status_code, 302)


class ReportsIndexTests(TestCase):
    def setUp(self):
        self.user = _user("superadmin")
        self.client = Client()
        self.client.force_login(self.user)

    def test_index_returns_200(self):
        resp = self.client.get(reverse("reports:index"))
        self.assertEqual(resp.status_code, 200)


class ReportsExcelTests(TestCase):
    def setUp(self):
        self.user = _user("superadmin")
        self.client = Client()
        self.client.force_login(self.user)

    def test_inventario_excel(self):
        resp = self.client.get(reverse("reports:inventario_excel"))
        self.assertEqual(resp.status_code, 200)
        self.assertIn("spreadsheetml", resp["Content-Type"])

    def test_usuarios_excel(self):
        resp = self.client.get(reverse("reports:usuarios_excel"))
        self.assertEqual(resp.status_code, 200)
        self.assertIn("spreadsheetml", resp["Content-Type"])

    def test_movimientos_excel(self):
        resp = self.client.get(reverse("reports:movimientos_excel"))
        self.assertEqual(resp.status_code, 200)
        self.assertIn("spreadsheetml", resp["Content-Type"])

    def test_costos_excel(self):
        resp = self.client.get(reverse("reports:costos_excel"))
        self.assertEqual(resp.status_code, 200)
        self.assertIn("spreadsheetml", resp["Content-Type"])

    def test_mantenimiento_excel(self):
        resp = self.client.get(reverse("reports:mantenimiento_excel"))
        self.assertEqual(resp.status_code, 200)
        self.assertIn("spreadsheetml", resp["Content-Type"])

    def test_excel_mantenimiento_columnas_nuevas(self):
        from io import BytesIO

        from openpyxl import load_workbook

        from mantenimiento.models import OrdenMantenimiento

        hoy = timezone.localdate()
        tecnico = CustomUser.objects.create_user(
            username="rpt_tecnico_excel", password="x", rol="tecnico", first_name="Ana", last_name="Técnica"
        )
        activo = Activo.objects.create(
            serial="SN-XLS-1", tipo_dispositivo="portatil", marca="HP", modelo="ProBook", estado="disponible"
        )
        OrdenMantenimiento.objects.create(
            activo=activo, tipo="preventivo", estado="cerrada", fecha_apertura=hoy, fecha_cierre=hoy,
            en_garantia=True, tecnico=tecnico, estado_final="operativo", costo_real=150000,
        )
        OrdenMantenimiento.objects.create(
            activo=activo, tipo="correctivo", estado="cerrada", fecha_apertura=hoy - timedelta(days=1),
            fecha_cierre=hoy, tecnico_asignado="Pedro Viejo", costo_estimado=80000,
        )

        resp = self.client.get(reverse("reports:mantenimiento_excel"))
        self.assertEqual(resp.status_code, 200)
        ws = load_workbook(BytesIO(resp.content)).active
        encabezados = [
            "Fecha apertura", "Fecha cierre", "Tipo", "En garantía", "Estado", "Estado final",
            "Prioridad", "Activo", "Serial", "Técnico", "Costo estimado", "Costo real",
        ]
        self.assertEqual([c.value for c in ws[6]], encabezados)
        preventivo = [c.value for c in ws[7]]
        correctivo = [c.value for c in ws[8]]
        self.assertEqual(preventivo[2:10], ["Preventivo", "Sí", "Cerrada", "Operativo", "Media", "HP ProBook", "SN-XLS-1", "Ana Técnica"])
        self.assertEqual(preventivo[11], 150000)
        self.assertEqual(ws.cell(row=7, column=12).number_format, "$#,##0.00")
        self.assertEqual(correctivo[3], "No")
        self.assertEqual(correctivo[9], "Pedro Viejo")
        self.assertEqual(correctivo[10], 80000)
        self.assertEqual(ws.cell(row=8, column=11).number_format, "$#,##0.00")


class ReportsPDFTests(TestCase):
    def setUp(self):
        self.user = _user("superadmin")
        self.client = Client()
        self.client.force_login(self.user)

    def test_inventario_pdf(self):
        with patch.dict(
            "sys.modules",
            {
                "weasyprint": _fake_weasyprint(),
                "weasyprint.HTML": _fake_weasyprint().HTML,
            },
        ):
            resp = self.client.get(reverse("reports:inventario_pdf"))
        self.assertEqual(resp.status_code, 200)

    def test_usuarios_pdf(self):
        with patch.dict(
            "sys.modules",
            {
                "weasyprint": _fake_weasyprint(),
                "weasyprint.HTML": _fake_weasyprint().HTML,
            },
        ):
            resp = self.client.get(reverse("reports:usuarios_pdf"))
        self.assertEqual(resp.status_code, 200)


class ReportsPermisosTests(TestCase):
    def test_tecnico_puede_ver_index(self):
        user = _user("tecnico")
        self.client.force_login(user)
        resp = self.client.get(reverse("reports:index"))
        self.assertEqual(resp.status_code, 200)


def _crear_usuario(nombre):
    return Usuario.objects.create(
        nombre_completo=nombre,
        documento_identidad=f"DOC-{Usuario.objects.count()}",
        cargo="Analista",
        area="Sistemas",
        correo=f"u{Usuario.objects.count()}@redihos.com",
    )


def _crear_activo(serial, **kwargs):
    defaults = {
        "tipo_dispositivo": "escritorio",
        "marca": "HP",
        "modelo": "EliteDesk",
        "estado": "asignado",
    }
    defaults.update(kwargs)
    return Activo.objects.create(serial=serial, **defaults)


class ReportsActivosPorUsuarioTests(TestCase):
    def setUp(self):
        self.user = _user("superadmin")
        self.client.force_login(self.user)
        self.url = reverse("reports:activos_por_usuario")

    def test_endpoint_requiere_login(self):
        self.client.logout()
        resp = self.client.get(self.url)
        self.assertEqual(resp.status_code, 302)

    def test_sin_q_muestra_placeholder(self):
        resp = self.client.get(self.url)
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, "Escribe un nombre o documento")

    def test_sin_resultados_muestra_placeholder(self):
        resp = self.client.get(self.url, {"q": "ZzZ inexistente"})
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, "Sin resultados")

    def test_con_resultado_muestra_tarjeta(self):
        persona = _crear_usuario("Maria Rodriguez")
        activo = _crear_activo("SN-ASIG-001")
        Asignacion.objects.create(
            activo=activo,
            usuario=persona,
            fecha_asignacion=timezone.now().date(),
            activa=True,
        )
        resp = self.client.get(self.url, {"q": "Maria"})
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, "Maria Rodriguez")
        self.assertContains(resp, "SN-ASIG-001")
        self.assertContains(resp, "1 activo asignado")


class ReportsLugarYGarantiaTests(TestCase):
    def setUp(self):
        self.user = _user("superadmin")
        self.client.force_login(self.user)
        self.url = reverse("reports:index")

    def test_index_lista_activos_por_lugar(self):
        _crear_activo("SN-LUG-001", ubicacion_fisica="Bodega Central")
        _crear_activo("SN-LUG-002", ubicacion_fisica="Bodega Central")
        resp = self.client.get(self.url)
        self.assertContains(resp, "Bodega Central")
        self.assertContains(resp, ">2<", html=False)

    def test_index_incluye_garantia_vencida(self):
        hoy = timezone.now().date()
        _crear_activo(
            "SN-GAR-001",
            fecha_compra=hoy - timedelta(days=400),
            garantia_fabrica_meses=12,
            proveedor="Distribuidora X",
        )
        resp = self.client.get(self.url)
        self.assertContains(resp, "SN-GAR-001")
        self.assertContains(resp, "Vencido")

    def test_index_incluye_garantia_vigente(self):
        hoy = timezone.now().date()
        _crear_activo(
            "SN-GAR-002",
            fecha_compra=hoy,
            garantia_fabrica_meses=24,
            proveedor="Distribuidora Y",
        )
        resp = self.client.get(self.url)
        self.assertContains(resp, "SN-GAR-002")
        self.assertContains(resp, "días")

    def test_index_incluye_garantia_extendida_sin_fabrica(self):
        hoy = timezone.now().date()
        _crear_activo(
            "SN-GAR-003",
            fecha_compra=hoy - timedelta(days=100),
            garantia_extendida=True,
            anios_garantia_extendida=2,
            proveedor="Distribuidora Z",
        )
        resp = self.client.get(self.url)
        self.assertContains(resp, "SN-GAR-003")
        self.assertContains(resp, "días")
