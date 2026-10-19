"""Vistas de reportes (PDF y Excel)."""
from datetime import timedelta
from decimal import Decimal
import re

from django.db.models import Count, Prefetch, Q
from django.http import HttpResponse
from django.shortcuts import render
from django.template.loader import render_to_string
from django.utils import timezone
from django.utils.html import format_html
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from accounts.permisos import requiere_permiso
from core.utils import guardar_excel

from inventario.models import Activo, Asignacion, Movimiento, TIPOS
from mantenimiento.models import OrdenMantenimiento
from usuarios.models import Usuario

from .datos import (
    CAMPOS_INVENTARIO,
    _build_costos,
    _build_report_context,
    _parse_date,
)
from .excel_estilo import (
    AZUL,
    AZUL_CLARO,
    AZUL_OSCURO,
    BLANCO,
    GRIS_TXT,
    LOGO_PATH,
    _adjust_column_widths,
    _apply_redihos_data_style,
    _apply_redihos_header,
    _style_summary_sheet,
)



# ── Vistas ───────────────────────────────────────────────────────────────────

@requiere_permiso("reportes", "lectura")
def index(request):
    today = timezone.now().date()
    rango_fin = _parse_date(request.GET.get("fin")) or today
    rango_inicio = _parse_date(request.GET.get("inicio")) or (rango_fin - timedelta(days=30))
    if rango_inicio > rango_fin:
        rango_inicio, rango_fin = rango_fin, rango_inicio

    context = _build_report_context(rango_inicio, rango_fin)
    return render(request, "reports/index.html", context)


@requiere_permiso("reportes", "lectura")
def activos_por_usuario(request):
    """Partial HTMX: una sola tarjeta con los activos actualmente asignados
    al usuario que coincide con la búsqueda. Sin búsqueda o sin coincidencia
    devuelve una tarjeta placeholder (nunca un error)."""
    q = request.GET.get("q", "").strip()
    usuario = None
    asignaciones = []

    if q:
        usuarios = list(
            Usuario.objects.filter(
                Q(nombre_completo__icontains=q)
                | Q(documento_identidad__icontains=q)
            ).order_by("nombre_completo")[:5]
        )
        if usuarios:
            usuario = usuarios[0]
            asignaciones = (
                Asignacion.objects.filter(usuario=usuario, activa=True)
                .select_related("activo")
                .order_by("-fecha_asignacion")
            )

    return render(request, "reports/partials/activos_usuario.html", {
        "q": q,
        "usuario": usuario,
        "asignaciones": asignaciones,
    })


@requiere_permiso("reportes", "lectura")
def inventario_pdf(request):
    import weasyprint

    # Selección de columnas
    campos_solicitados = request.GET.getlist("campos")
    campos = [c for c in campos_solicitados if c in CAMPOS_INVENTARIO] or list(CAMPOS_INVENTARIO.keys())

    # Filtro por tipo de dispositivo
    tipos_solicitados = request.GET.getlist("tipos")
    filtrar_por_tipo = tipos_solicitados and "todos" not in tipos_solicitados

    activos_qs = Activo.objects.prefetch_related(
        Prefetch(
            "asignaciones",
            queryset=Asignacion.objects.filter(activa=True).select_related("usuario"),
            to_attr="asignaciones_activas",
        )
    ).order_by("tipo_dispositivo", "marca", "modelo")

    if filtrar_por_tipo:
        activos_qs = activos_qs.filter(tipo_dispositivo__in=tipos_solicitados)

    # Construir filas: cada fila es un activo, cada celda un campo
    filas = []
    for activo in activos_qs:
        asignacion = activo.asignaciones_activas[0] if activo.asignaciones_activas else None
        fila = []
        for c in campos:
            raw = CAMPOS_INVENTARIO[c][1](activo, asignacion)
            if c == "estado":
                valor = format_html('<span class="badge {}">{}</span>', activo.estado, raw)
                vacio = False
            elif c == "valor_compra" and raw is not None:
                valor = f"${raw:,.0f} COP"
                vacio = False
            else:
                vacio = raw in (None, "")
                valor = "—" if vacio else raw
            fila.append({"valor": valor, "vacio": vacio})
        filas.append(fila)

    # Podar columnas sin datos en ninguna fila (patrón de Snipe-IT/GLPI:
    # un campo vacío para todo el universo no aporta y ensancha el reporte)
    if filas:
        keep_idx = [
            i for i in range(len(campos))
            if any(not fila[i]["vacio"] for fila in filas)
        ]
        if len(keep_idx) < len(campos):
            campos = [campos[i] for i in keep_idx]
            filas = [[fila[i] for i in keep_idx] for fila in filas]

    columnas_count = len(campos)
    if columnas_count > 25:
        tabla_font = "6pt"
    elif columnas_count > 15:
        tabla_font = "7pt"
    else:
        tabla_font = "8.5pt"

    # Anchos dinámicos según contenido (estilo GLPI Protocols Manager:
    # "dynamic column widths"). El peso usa el token más largo (no la suma
    # de caracteres), porque las celdas envuelven por palabra. Se normaliza
    # a 100% con table-layout: fixed para que nada se salga de la página.
    TOKEN_CAPS = {
        "id": 6, "estado": 15, "valor_compra": 15, "observaciones": 26,
        "fecha_compra": 11, "fecha_vencimiento_garantia": 11,
        "fecha_asignacion": 11, "fecha_creacion": 17,
        "ultima_actualizacion": 17, "documento_usuario": 12,
    }

    def _peso_columna(campo, idx):
        tokens = [t for t in re.split(r"\s+", CAMPOS_INVENTARIO[campo][0]) if t]
        for fila in filas:
            texto = re.sub(r"<[^>]+>", "", str(fila[idx]["valor"]))
            tokens.extend(t for t in re.split(r"\s+", texto) if t)
        mayor = max((len(t) for t in tokens), default=4)
        cap = TOKEN_CAPS.get(campo, 22)
        return min(max(mayor + 1, 6), cap)

    pesos = [_peso_columna(c, idx) for idx, c in enumerate(campos)]
    total_peso = sum(pesos) or 1
    columnas = [
        {
            "key": c,
            "label": CAMPOS_INVENTARIO[c][0],
            "ancho": round(pesos[idx] / total_peso * 100, 2),
        }
        for idx, c in enumerate(campos)
    ]

    tipos_a_resumir = (
        [(key, label) for key, label in TIPOS if key in tipos_solicitados]
        if filtrar_por_tipo else TIPOS
    )
    resumen = []
    for key, label in tipos_a_resumir:
        qs = activos_qs.filter(tipo_dispositivo=key)
        resumen.append({
            "label": label,
            "total": qs.count(),
            "disponibles": qs.filter(estado="disponible").count(),
            "asignados": qs.filter(estado="asignado").count(),
            "mantenimiento": qs.filter(estado="en_mantenimiento").count(),
            "reparacion": qs.filter(estado="en_reparacion").count(),
            "baja": qs.filter(estado="dado_de_baja").count(),
        })

    totales = {"activos": activos_qs.count()}
    if "valor_compra" in campos:
        total_valor = sum(
            (a.valor_compra or Decimal("0")) for a in activos_qs.exclude(valor_compra__isnull=True)
        )
        totales["valor"] = f"${total_valor:,.0f} COP"
    if "estado" in campos:
        totales["estado"] = {
            "disponibles": activos_qs.filter(estado="disponible").count(),
            "asignados": activos_qs.filter(estado="asignado").count(),
            "mantenimiento": activos_qs.filter(estado="en_mantenimiento").count(),
            "reparacion": activos_qs.filter(estado="en_reparacion").count(),
            "baja": activos_qs.filter(estado="dado_de_baja").count(),
        }

    html_string = render_to_string("reports/inventario_pdf.html", {
        "columnas": columnas,
        "filas": filas,
        "resumen": resumen,
        "totales": totales,
        "campos": campos,
        "tabla_font": tabla_font,
        "logo_path": LOGO_PATH,
        "generado": timezone.now(),
    })
    pdf_bytes = weasyprint.HTML(
        string=html_string, base_url=request.build_absolute_uri("/")
    ).write_pdf()

    response = HttpResponse(pdf_bytes, content_type="application/pdf")
    response["Content-Disposition"] = 'inline; filename="reporte_inventario.pdf"'
    return response


@requiere_permiso("reportes", "lectura")
def usuarios_pdf(request):
    import weasyprint

    usuarios = (
        Usuario.objects.annotate(
            activos_asignados=Count("asignaciones", filter=Q(asignaciones__activa=True))
        )
        .order_by("nombre_completo")
    )
    context = {
        "fecha_generacion": timezone.now(),
        "usuarios": usuarios,
    }

    html_string = render_to_string("reports/usuarios_pdf.html", context)
    pdf_bytes = weasyprint.HTML(
        string=html_string, base_url=request.build_absolute_uri("/")
    ).write_pdf()

    response = HttpResponse(pdf_bytes, content_type="application/pdf")
    response["Content-Disposition"] = 'inline; filename="reporte_usuarios.pdf"'
    return response


@requiere_permiso("reportes", "lectura")
def inventario_excel(request):
    wb = Workbook()
    ws = wb.active
    ws.title = "Inventario"

    # Selección de columnas
    campos_solicitados = request.GET.getlist("campos")
    if campos_solicitados:
        campos = [c for c in campos_solicitados if c in CAMPOS_INVENTARIO]
    else:
        campos = list(CAMPOS_INVENTARIO.keys())

    # Filtro por tipo de dispositivo
    tipos_solicitados = request.GET.getlist("tipos")
    filtrar_por_tipo = tipos_solicitados and "todos" not in tipos_solicitados
    tipos_a_resumir = (
        [(key, label) for key, label in TIPOS if key in tipos_solicitados]
        if filtrar_por_tipo else TIPOS
    )

    headers = [CAMPOS_INVENTARIO[c][0] for c in campos]
    last_col_letter = get_column_letter(len(headers))
    start_row = _apply_redihos_header(ws, headers, "Inventario")

    activos = Activo.objects.prefetch_related(
        Prefetch(
            "asignaciones",
            queryset=Asignacion.objects.filter(activa=True).select_related("usuario"),
            to_attr="asignaciones_activas",
        )
    ).order_by("tipo_dispositivo", "marca", "modelo")

    if filtrar_por_tipo:
        activos = activos.filter(tipo_dispositivo__in=tipos_solicitados)

    for _row_idx, activo in enumerate(activos, start=start_row):
        asignacion = activo.asignaciones_activas[0] if activo.asignaciones_activas else None
        row = [CAMPOS_INVENTARIO[c][1](activo, asignacion) for c in campos]
        ws.append(row)

    _apply_redihos_data_style(ws, start_row, len(headers))

    # Formato monetario para valor_compra si está presente
    if "valor_compra" in campos:
        valor_col = campos.index("valor_compra") + 1
        for row_idx in range(start_row, ws.max_row + 1):
            ws.cell(row=row_idx, column=valor_col).number_format = '$#,##0'

    # ── Resumen por categorías (tipo de dispositivo) ─────────────────────────
    thin_side = Side(style="thin", color="D1D5DB")
    border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
    data_font = Font(name="Calibri", size=10, color=GRIS_TXT)
    white_fill = PatternFill(start_color=BLANCO, end_color=BLANCO, fill_type="solid")
    blue_light_fill = PatternFill(start_color=AZUL_CLARO, end_color=AZUL_CLARO, fill_type="solid")
    section_title_font = Font(name="Calibri", bold=True, size=12, color=AZUL_OSCURO)
    section_title_fill = PatternFill(start_color=AZUL_CLARO, end_color=AZUL_CLARO, fill_type="solid")
    header_fill = PatternFill(start_color=AZUL, end_color=AZUL, fill_type="solid")
    header_font = Font(name="Calibri", bold=True, size=10, color=BLANCO)

    resumen_start = ws.max_row + 3
    ws.merge_cells(f"A{resumen_start}:{last_col_letter}{resumen_start}")
    title_cell = ws.cell(row=resumen_start, column=1, value="Resumen por categorías")
    title_cell.font = section_title_font
    title_cell.fill = section_title_fill
    title_cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[resumen_start].height = 22

    resumen_headers = [
        "Categoría", "Total", "Disponibles", "Asignados",
        "En mantenimiento", "En reparación", "Dados de baja"
    ]
    header_row = resumen_start + 1
    for col_idx, h in enumerate(resumen_headers, start=1):
        cell = ws.cell(row=header_row, column=col_idx, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
    ws.row_dimensions[header_row].height = 26

    data_start = header_row + 1
    for idx, (key, label) in enumerate(tipos_a_resumir, start=data_start):
        qs = activos.filter(tipo_dispositivo=key)
        values = [
            label,
            qs.count(),
            qs.filter(estado="disponible").count(),
            qs.filter(estado="asignado").count(),
            qs.filter(estado="en_mantenimiento").count(),
            qs.filter(estado="en_reparacion").count(),
            qs.filter(estado="dado_de_baja").count(),
        ]
        fill = white_fill if (idx - data_start) % 2 == 0 else blue_light_fill
        for col_idx, val in enumerate(values, start=1):
            cell = ws.cell(row=idx, column=col_idx, value=val)
            cell.font = data_font
            cell.fill = fill
            cell.border = border
            cell.alignment = Alignment(
                horizontal="left" if col_idx == 1 else "center",
                vertical="center",
                indent=1
            )

    # ── Fila de totales generales ────────────────────────────────────────────
    totales_font = Font(name="Calibri", bold=True, size=11, color=AZUL_OSCURO)
    totales_fill = PatternFill(start_color=AZUL_CLARO, end_color=AZUL_CLARO, fill_type="solid")
    total_row = ws.max_row + 2

    ws.merge_cells(f"A{total_row}:{last_col_letter}{total_row}")
    total_cell = ws.cell(row=total_row, column=1, value=f"Total de activos: {activos.count()}")
    total_cell.font = totales_font
    total_cell.fill = totales_fill
    total_cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[total_row].height = 22

    current_row = total_row + 1

    if "valor_compra" in campos:
        total_valor = sum(
            (a.valor_compra or Decimal("0")) for a in activos.exclude(valor_compra__isnull=True)
        )
        ws.merge_cells(f"A{current_row}:{last_col_letter}{current_row}")
        valor_cell = ws.cell(
            row=current_row, column=1,
            value=f"Valor total inventario: ${total_valor:,.0f} COP"
        )
        valor_cell.font = totales_font
        valor_cell.fill = totales_fill
        valor_cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        ws.row_dimensions[current_row].height = 22
        current_row += 1

    if "estado" in campos:
        qs = activos
        ws.merge_cells(f"A{current_row}:{last_col_letter}{current_row}")
        estado_cell = ws.cell(
            row=current_row, column=1,
            value=(
                f"Por estado — Disponibles: {qs.filter(estado='disponible').count()} · "
                f"Asignados: {qs.filter(estado='asignado').count()} · "
                f"En mantenimiento: {qs.filter(estado='en_mantenimiento').count()} · "
                f"En reparación: {qs.filter(estado='en_reparacion').count()} · "
                f"Dados de baja: {qs.filter(estado='dado_de_baja').count()}"
            )
        )
        estado_cell.font = totales_font
        estado_cell.fill = totales_fill
        estado_cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        ws.row_dimensions[current_row].height = 22

    data_end_row = start_row + activos.count() - 1
    _adjust_column_widths(ws, headers, data_end_row=data_end_row)

    # Asegurar anchos mínimos para las columnas del resumen
    for col_idx, h in enumerate(resumen_headers, start=1):
        col_letter = get_column_letter(col_idx)
        current = ws.column_dimensions[col_letter].width or 12
        ws.column_dimensions[col_letter].width = max(current, len(h) + 4)

    # Configuración de página: horizontal, ajustar ancho y centrado horizontal
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True
    ws.print_options.verticalCentered = False

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = "attachment; filename=reporte_inventario.xlsx"
    guardar_excel(wb, response)
    return response


@requiere_permiso("reportes", "lectura")
def usuarios_excel(request):
    wb = Workbook()

    usuarios_activos = Usuario.objects.filter(estado="activo").count()
    usuarios_inactivos = Usuario.objects.filter(estado="inactivo").count()
    usuarios_por_area = (
        Usuario.objects.values("area")
        .annotate(total=Count("id"))
        .order_by("-total", "area")
    )
    usuarios_por_cargo = (
        Usuario.objects.values("cargo")
        .annotate(total=Count("id"))
        .order_by("-total", "cargo")
    )

    ws_resumen = wb.active
    ws_resumen.title = "Resumen"
    _style_summary_sheet(ws_resumen, "Usuarios")

    ws_resumen["A6"] = "Usuarios activos"
    ws_resumen["B6"] = usuarios_activos
    ws_resumen["A7"] = "Usuarios inactivos"
    ws_resumen["B7"] = usuarios_inactivos
    ws_resumen["A9"] = "Por area"
    ws_resumen["B9"] = "Total"
    for idx, row in enumerate(usuarios_por_area, start=10):
        ws_resumen.cell(row=idx, column=1, value=row["area"])
        ws_resumen.cell(row=idx, column=2, value=row["total"])
    offset = 10 + len(usuarios_por_area) + 2
    ws_resumen.cell(row=offset, column=1, value="Por cargo")
    ws_resumen.cell(row=offset, column=2, value="Total")
    for idx, row in enumerate(usuarios_por_cargo, start=offset + 1):
        ws_resumen.cell(row=idx, column=1, value=row["cargo"])
        ws_resumen.cell(row=idx, column=2, value=row["total"])

    ws = wb.create_sheet(title="Usuarios")
    headers = [
        "Nombre", "Documento", "Cargo", "Area", "Correo",
        "Telefono", "Estado", "Fecha creacion"
    ]
    start_row = _apply_redihos_header(ws, headers, "Usuarios")

    for usuario in Usuario.objects.all().order_by("nombre_completo"):
        ws.append([
            usuario.nombre_completo,
            usuario.documento_identidad,
            usuario.cargo,
            usuario.area,
            usuario.correo,
            usuario.telefono,
            usuario.get_estado_display(),
            usuario.fecha_creacion.strftime("%Y-%m-%d"),
        ])

    _apply_redihos_data_style(ws, start_row, len(headers))
    _adjust_column_widths(ws, headers)

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = "attachment; filename=reporte_usuarios.xlsx"
    guardar_excel(wb, response)
    return response


@requiere_permiso("reportes", "lectura")
def costos_excel(request):
    today = timezone.now().date()
    costos = _build_costos(today)

    wb = Workbook()
    ws_resumen = wb.active
    ws_resumen.title = "Resumen"
    _style_summary_sheet(ws_resumen, "Costos")

    ws_resumen["A6"] = "Vida util (meses)"
    ws_resumen["B6"] = costos["vida_util"]
    ws_resumen["A7"] = "Valor compra total"
    ws_resumen["B7"] = float(costos["total"])
    ws_resumen["A8"] = "Depreciacion acumulada"
    ws_resumen["B8"] = float(costos["depreciacion"])
    ws_resumen["A9"] = "Valor actual estimado"
    ws_resumen["B9"] = float(costos["valor_actual"])

    ws = wb.create_sheet(title="Costos")
    headers = [
        "Activo", "Serial", "Tipo", "Valor compra",
        "Meses", "Depreciacion", "Valor actual"
    ]
    start_row = _apply_redihos_header(ws, headers, "Costos")

    for item in costos["items"]:
        activo = item["activo"]
        ws.append([
            str(activo),
            activo.serial,
            activo.get_tipo_dispositivo_display(),
            float(item["valor"]),
            item["meses"],
            float(item["depreciacion"]),
            float(item["valor_actual"]),
        ])

    _apply_redihos_data_style(ws, start_row, len(headers))
    money_cols = {4, 6, 7}  # columnas con valores monetarios
    for row_idx in range(start_row, ws.max_row + 1):
        for col in money_cols:
            ws.cell(row=row_idx, column=col).number_format = '$#,##0.00'
    _adjust_column_widths(ws, headers)

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = "attachment; filename=reporte_costos.xlsx"
    guardar_excel(wb, response)
    return response


@requiere_permiso("reportes", "lectura")
def movimientos_excel(request):
    today = timezone.now().date()
    rango_fin = _parse_date(request.GET.get("fin")) or today
    rango_inicio = _parse_date(request.GET.get("inicio")) or (rango_fin - timedelta(days=30))
    if rango_inicio > rango_fin:
        rango_inicio, rango_fin = rango_fin, rango_inicio

    movimientos = (
        Movimiento.objects.filter(fecha__range=(rango_inicio, rango_fin))
        .select_related("activo", "usuario_destino")
        .order_by("-fecha", "-fecha_creacion")
    )

    wb = Workbook()
    ws = wb.active
    ws.title = "Movimientos"
    headers = [
        "Fecha", "Tipo", "Activo", "Serial", "Usuario destino",
        "Descripcion", "Realizado por"
    ]
    start_row = _apply_redihos_header(ws, headers, "Movimientos")

    for mov in movimientos:
        ws.append([
            mov.fecha.strftime("%Y-%m-%d"),
            mov.get_tipo_display(),
            mov.activo.get_tipo_dispositivo_display(),
            mov.activo.serial,
            mov.usuario_destino.nombre_completo if mov.usuario_destino else "",
            mov.descripcion,
            mov.realizado_por,
        ])

    _apply_redihos_data_style(ws, start_row, len(headers))
    _adjust_column_widths(ws, headers)

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = "attachment; filename=reporte_movimientos.xlsx"
    guardar_excel(wb, response)
    return response


@requiere_permiso("reportes", "lectura")
def mantenimiento_excel(request):
    today = timezone.now().date()
    rango_fin = _parse_date(request.GET.get("fin")) or today
    rango_inicio = _parse_date(request.GET.get("inicio")) or (rango_fin - timedelta(days=30))
    if rango_inicio > rango_fin:
        rango_inicio, rango_fin = rango_fin, rango_inicio

    ordenes = (
        OrdenMantenimiento.objects.filter(fecha_apertura__range=(rango_inicio, rango_fin))
        .select_related("activo", "plan")
        .order_by("-fecha_apertura", "-fecha_creacion")
    )

    wb = Workbook()
    ws = wb.active
    ws.title = "Ordenes"
    headers = [
        "Fecha apertura",
        "Fecha cierre",
        "Tipo",
        "Estado",
        "Prioridad",
        "Activo",
        "Serial",
        "Plan",
        "Tecnico",
        "Costo estimado",
        "Costo real",
    ]
    start_row = _apply_redihos_header(ws, headers, "Mantenimiento")

    for orden in ordenes:
        ws.append([
            orden.fecha_apertura.strftime("%Y-%m-%d") if orden.fecha_apertura else "",
            orden.fecha_cierre.strftime("%Y-%m-%d") if orden.fecha_cierre else "",
            orden.get_tipo_display(),
            orden.get_estado_display(),
            orden.get_prioridad_display(),
            f"{orden.activo.marca} {orden.activo.modelo}",
            orden.activo.serial,
            f"Plan #{orden.plan_id}" if orden.plan_id else "",
            orden.tecnico_asignado,
            float(orden.costo_estimado) if orden.costo_estimado else "",
            float(orden.costo_real) if orden.costo_real else "",
        ])

    _apply_redihos_data_style(ws, start_row, len(headers))
    money_cols = {10, 11}
    for row_idx in range(start_row, ws.max_row + 1):
        for col in money_cols:
            cell = ws.cell(row=row_idx, column=col)
            if isinstance(cell.value, (int, float)):
                cell.number_format = '$#,##0.00'
    _adjust_column_widths(ws, headers)

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = "attachment; filename=reporte_mantenimiento.xlsx"
    guardar_excel(wb, response)
    return response
