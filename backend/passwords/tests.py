import json
import threading
import zipfile
from contextlib import contextmanager
from datetime import date, time, timedelta
from io import BytesIO, StringIO
from unittest import skipUnless
from unittest.mock import patch

from django.conf import settings
from django.contrib.auth import get_user_model
from django.db import connection
from django.test import Client, TestCase, TransactionTestCase
from django.urls import reverse
from openpyxl import Workbook

from passwords.crypto import build_fernet
from passwords.models import AccesoLog, Credencial, Vault

CustomUser = get_user_model()

ENCABEZADOS = ("Título", "Usuario", "Contraseña", "URL", "Notas", "Estado", "Fecha de expiración")


def _xlsx(filas, encabezados=ENCABEZADOS, nombre="claves.xlsx"):
    wb = Workbook()
    ws = wb.active
    ws.append(list(encabezados))
    for fila in filas:
        ws.append(list(fila))
    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    buffer.name = nombre
    return buffer


def _xlsx_a_mano(filas_xml="", shared_strings=None, hojas_declaradas=1, partes_extra=None):
    """xlsx armado a mano, como un archivo fabricado (openpyxl no deja escribir estos casos).

    La hoja trae los encabezados Título y Contraseña; filas_xml va después, dentro de
    <sheetData>. Todas las hojas declaradas apuntan a la misma parte.
    """
    ns = 'xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"'
    rel = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    tipo = "application/vnd.openxmlformats-officedocument.spreadsheetml"
    rels_ns = 'xmlns="http://schemas.openxmlformats.org/package/2006/relationships"'
    tipos = (
        '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
        '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
        '<Default Extension="xml" ContentType="application/xml"/>'
        f'<Override PartName="/xl/workbook.xml" ContentType="{tipo}.sheet.main+xml"/>'
        f'<Override PartName="/xl/worksheets/sheet1.xml" ContentType="{tipo}.worksheet+xml"/>'
    )
    rels_libro = f'<Relationship Id="rId1" Type="{rel}/worksheet" Target="worksheets/sheet1.xml"/>'
    if shared_strings is not None:
        tipos += f'<Override PartName="/xl/sharedStrings.xml" ContentType="{tipo}.sharedStrings+xml"/>'
        rels_libro += f'<Relationship Id="rId2" Type="{rel}/sharedStrings" Target="sharedStrings.xml"/>'
    hojas = "".join(f'<sheet name="H{i}" sheetId="{i + 1}" r:id="rId1"/>' for i in range(hojas_declaradas))
    partes = {
        "[Content_Types].xml": tipos + "</Types>",
        "_rels/.rels": f'<Relationships {rels_ns}><Relationship Id="rId1" Type="{rel}/officeDocument" '
                       'Target="xl/workbook.xml"/></Relationships>',
        "xl/workbook.xml": f'<workbook {ns} xmlns:r="{rel}"><sheets>{hojas}</sheets></workbook>',
        "xl/_rels/workbook.xml.rels": f"<Relationships {rels_ns}>{rels_libro}</Relationships>",
        "xl/worksheets/sheet1.xml": (
            f'<worksheet {ns}><sheetData><row r="1">'
            '<c r="A1" t="inlineStr"><is><t>Título</t></is></c>'
            '<c r="B1" t="inlineStr"><is><t>Contraseña</t></is></c>'
            f"</row>{filas_xml}</sheetData></worksheet>"
        ),
    }
    if shared_strings is not None:
        partes["xl/sharedStrings.xml"] = f"<sst {ns}>{shared_strings}</sst>"
    partes.update(partes_extra or {})
    buffer = BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as paquete:
        for nombre, contenido in partes.items():
            paquete.writestr(nombre, contenido)
    buffer.seek(0)
    buffer.name = "fabricado.xlsx"
    return buffer


@contextmanager
def _lectura_acotada(max_filas, max_ancho):
    """Cuenta las filas que openpyxl entrega al leer y falla en cuanto se pasa.

    Mide el trabajo sin reloj: si la lectura deja de estar acotada, la prueba falla
    en el acto en vez de quedarse minutos recorriendo filas vacías.
    """
    from openpyxl.worksheet._read_only import ReadOnlyWorksheet

    original = ReadOnlyWorksheet._cells_by_row
    medida = {"filas": 0}

    def contando(ws, *args, **kwargs):
        for fila in original(ws, *args, **kwargs):
            medida["filas"] += 1
            if medida["filas"] > max_filas or len(fila) > max_ancho:
                raise AssertionError(f"Lectura sin límite: {medida['filas']} filas, ancho {len(fila)}")
            yield fila

    with patch.object(ReadOnlyWorksheet, "_cells_by_row", contando):
        yield medida


def _descifrar(token):
    return build_fernet().decrypt(token.encode()).decode()


def _user(rol):
    return CustomUser.objects.create_user(
        username=f"pw_{rol}_{CustomUser.objects.count()}",
        password="x",
        rol=rol,
    )


def _vault(nombre="Vault Test", **kw):
    return Vault.objects.create(nombre=nombre, **kw)


# ─── Vault tests ───────────────────────────────────────────────────


class VaultLoginTests(TestCase):
    def test_index_requiere_login(self):
        resp = self.client.get(reverse("passwords:index"))
        self.assertEqual(resp.status_code, 302)


class VaultCRUDTests(TestCase):
    def setUp(self):
        self.user = _user("superadmin")
        self.client = Client()
        self.client.force_login(self.user)

    def test_index_returns_200(self):
        resp = self.client.get(reverse("passwords:index"))
        self.assertEqual(resp.status_code, 200)

    def test_crear_vault(self):
        resp = self.client.post(
            reverse("passwords:vault_crear"),
            {"nombre": "Nuevo Vault", "descripcion": "Test"},
        )
        self.assertEqual(resp.status_code, 302)
        self.assertTrue(Vault.objects.filter(nombre="Nuevo Vault").exists())

    def test_editar_vault(self):
        v = _vault()
        resp = self.client.post(
            reverse("passwords:vault_editar", args=[v.pk]),
            {"nombre": "Editado", "descripcion": "x"},
        )
        self.assertEqual(resp.status_code, 302)
        v.refresh_from_db()
        self.assertEqual(v.nombre, "Editado")

    def test_eliminar_vault(self):
        v = _vault()
        resp = self.client.post(reverse("passwords:vault_eliminar", args=[v.pk]))
        self.assertEqual(resp.status_code, 302)
        self.assertFalse(Vault.objects.filter(pk=v.pk).exists())

    def test_la_pagina_de_confirmar_borrado_del_vault_abre(self):
        """Daba 500 en produccion: la plantilla usaba {% url x **kwargs %}, que Django no soporta."""
        v = _vault()
        resp = self.client.get(reverse("passwords:vault_eliminar", args=[v.pk]))
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, "Eliminar vault")
        self.assertContains(resp, v.nombre)
        self.assertContains(resp, f'href="{reverse("passwords:index")}"')
        self.assertTrue(Vault.objects.filter(pk=v.pk).exists())

    def test_la_pagina_de_confirmar_borrado_de_una_credencial_abre(self):
        v = _vault()
        cred = Credencial.objects.create(vault=v, titulo="T", usuario="u")
        resp = self.client.get(reverse("passwords:credencial_eliminar", args=[v.pk, cred.pk]))
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, f'href="{reverse("passwords:credenciales_lista", args=[v.pk])}"')

    def test_vault_crear_log(self):
        self.client.post(
            reverse("passwords:vault_crear"),
            {"nombre": "Con Log"},
        )
        self.assertTrue(AccesoLog.objects.filter(accion="vault_creado").exists())


# ─── Credencial tests ─────────────────────────────────────────────


class CredencialCRUDTests(TestCase):
    def setUp(self):
        self.user = _user("superadmin")
        self.client = Client()
        self.client.force_login(self.user)
        self.vault = _vault()

    def test_lista_credenciales(self):
        resp = self.client.get(
            reverse("passwords:credenciales_lista", args=[self.vault.pk])
        )
        self.assertEqual(resp.status_code, 200)

    def test_crear_credencial_con_secreto(self):
        resp = self.client.post(
            reverse("passwords:credencial_crear", args=[self.vault.pk]),
            {
                "titulo": "AWS Root",
                "usuario": "admin",
                "secreto": "mi_password_123",
                "estado": "activa",
            },
        )
        self.assertEqual(resp.status_code, 302)
        cred = Credencial.objects.get(titulo="AWS Root")
        self.assertNotEqual(cred.secreto_encriptado, "")
        self.assertEqual(cred.get_secret(), "mi_password_123")

    def test_editar_credencial_sin_cambiar_secreto(self):
        cred = Credencial.objects.create(
            vault=self.vault,
            titulo="Original",
            creado_por=self.user,
        )
        cred.set_secret("old_secret")
        cred.save()
        resp = self.client.post(
            reverse("passwords:credencial_editar", args=[self.vault.pk, cred.pk]),
            {
                "titulo": "Cambiado",
                "usuario": "",
                "secreto": "",
                "estado": "activa",
            },
        )
        self.assertEqual(resp.status_code, 302)
        cred.refresh_from_db()
        self.assertEqual(cred.titulo, "Cambiado")
        self.assertEqual(cred.get_secret(), "old_secret")

    def test_eliminar_credencial(self):
        cred = Credencial.objects.create(
            vault=self.vault, titulo="Del", creado_por=self.user
        )
        resp = self.client.post(
            reverse("passwords:credencial_eliminar", args=[self.vault.pk, cred.pk])
        )
        self.assertEqual(resp.status_code, 302)
        self.assertFalse(Credencial.objects.filter(pk=cred.pk).exists())


# ─── Acceso код / cifrado ──────────────────────────────────────────


class VaultAccessCodeTests(TestCase):
    def setUp(self):
        self.user = _user("superadmin")
        self.client = Client()
        self.client.force_login(self.user)

    def test_vault_con_codigo_acceso(self):
        v = _vault(nombre="Secreto", acceso_requerido=True)
        v.set_access_code("1234")
        v.save()
        self.assertTrue(v.check_access_code("1234"))
        self.assertFalse(v.check_access_code("wrong"))

    def test_vault_sin_codigo_acceso(self):
        v = _vault(nombre="Abierto", acceso_requerido=False)
        self.assertTrue(v.check_access_code(""))

    def test_ver_secreto_con_codigo_correcto(self):
        v = _vault(nombre="Locked", acceso_requerido=True)
        v.set_access_code("5678")
        v.save()
        cred = Credencial.objects.create(vault=v, titulo="C1")
        cred.set_secret("top_secret")
        cred.save()
        resp = self.client.post(
            reverse("passwords:credencial_secreto", args=[v.pk, cred.pk]),
            {"acceso_codigo": "5678"},
        )
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp.context["secreto"], "top_secret")

    def test_ver_secreto_con_codigo_incorrecto(self):
        v = _vault(nombre="Locked2", acceso_requerido=True)
        v.set_access_code("0000")
        v.save()
        cred = Credencial.objects.create(vault=v, titulo="C2")
        cred.set_secret("secret2")
        cred.save()
        resp = self.client.post(
            reverse("passwords:credencial_secreto", args=[v.pk, cred.pk]),
            {"acceso_codigo": "wrong"},
        )
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp.context["secreto"], "")


# ─── Export + Logs ─────────────────────────────────────────────────


class PasswordsExportTests(TestCase):
    def setUp(self):
        self.user = _user("superadmin")
        self.client = Client()
        self.client.force_login(self.user)
        self.vault = _vault()

    def test_export_excel(self):
        resp = self.client.get(
            reverse("passwords:credenciales_export", args=[self.vault.pk])
        )
        self.assertEqual(resp.status_code, 200)
        self.assertIn("spreadsheetml", resp["Content-Type"])

    def test_logs_returns_200(self):
        resp = self.client.get(reverse("passwords:logs", args=[self.vault.pk]))
        self.assertEqual(resp.status_code, 200)


class PasswordsPermisosTests(TestCase):
    def test_lectura_no_puede_crear_vault(self):
        user = _user("lectura")
        self.client.force_login(user)
        resp = self.client.get(reverse("passwords:vault_crear"))
        self.assertEqual(resp.status_code, 302)


class SecretoPermisosTests(TestCase):
    """Ver secretos en claro: solo superadmin/admin, con límite de intentos."""

    def setUp(self):
        from django.core.cache import cache

        cache.clear()
        self.vault = _vault(nombre="Restringido", acceso_requerido=True)
        self.vault.set_access_code("1111")
        self.vault.save()
        self.cred = Credencial.objects.create(vault=self.vault, titulo="C")
        self.cred.set_secret("s3creto")
        self.cred.save()
        self.url = reverse("passwords:credencial_secreto", args=[self.vault.pk, self.cred.pk])

    def test_rol_lectura_no_puede_ver_secreto(self):
        self.client.force_login(_user("lectura"))
        resp = self.client.post(self.url, {"acceso_codigo": "1111"})
        self.assertRedirects(
            resp, reverse("passwords:credenciales_lista", args=[self.vault.pk])
        )
        self.assertFalse(AccesoLog.objects.filter(accion="secreto_visto").exists())

    def test_rol_lectura_no_ve_boton_de_secreto(self):
        self.client.force_login(_user("lectura"))
        resp = self.client.get(reverse("passwords:credenciales_lista", args=[self.vault.pk]))
        self.assertEqual(resp.status_code, 200)
        self.assertNotContains(resp, self.url)

    def test_admin_puede_ver_secreto(self):
        self.client.force_login(_user("admin"))
        resp = self.client.post(self.url, {"acceso_codigo": "1111"})
        self.assertEqual(resp.context["secreto"], "s3creto")

    def test_codigo_invalido_queda_registrado(self):
        self.client.force_login(_user("superadmin"))
        self.client.post(self.url, {"acceso_codigo": "mal"})
        self.assertTrue(AccesoLog.objects.filter(accion="codigo_invalido").exists())

    def test_bloqueo_tras_intentos_fallidos(self):
        from passwords.views import MAX_INTENTOS_CODIGO

        self.client.force_login(_user("superadmin"))
        for _ in range(MAX_INTENTOS_CODIGO):
            self.client.post(self.url, {"acceso_codigo": "mal"})
        resp = self.client.post(self.url, {"acceso_codigo": "1111"})
        self.assertEqual(resp.context["secreto"], "")


class RotarClaveVaultTests(TestCase):
    def test_rota_secretos_legados_a_la_clave_nueva(self):
        from cryptography.fernet import Fernet
        from django.core.management import call_command
        from django.test import override_settings

        vault = _vault()
        cred = Credencial(vault=vault, titulo="Legado")
        with override_settings(PASSWORDS_ENCRYPTION_KEY=""):
            cred.set_secret("viejo")
            cred.save()

        clave = Fernet.generate_key().decode()
        with override_settings(PASSWORDS_ENCRYPTION_KEY=clave):
            # Con la clave nueva todavía se lee lo legado.
            self.assertEqual(Credencial.objects.get(pk=cred.pk).get_secret(), "viejo")
            call_command("rotar_clave_vault", stdout=StringIO())
            token = Credencial.objects.get(pk=cred.pk).secreto_encriptado
            self.assertEqual(Fernet(clave).decrypt(token.encode()).decode(), "viejo")


class SecretoIlegibleTests(TestCase):
    def test_secreto_con_clave_desconocida_no_da_500(self):
        from cryptography.fernet import Fernet

        vault = _vault()
        cred = Credencial.objects.create(
            vault=vault,
            titulo="Vieja",
            secreto_encriptado=Fernet(Fernet.generate_key()).encrypt(b"x").decode(),
        )
        self.client.force_login(_user("superadmin"))
        resp = self.client.post(
            reverse("passwords:credencial_secreto", args=[vault.pk, cred.pk]), {}
        )
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp.context["secreto"], "")
        self.assertTrue(AccesoLog.objects.filter(accion="secreto_ilegible").exists())


# ─── Importación desde Excel: lectura y validación ───────────────


class ImportacionLecturaTests(TestCase):
    def setUp(self):
        self.vault = _vault()

    def _leer(self, filas, **kw):
        from passwords.importacion import leer_credenciales
        return leer_credenciales(_xlsx(filas, **kw), self.vault)[:2]

    def test_fila_valida_queda_nueva_y_cifrada(self):
        filas, avisos = self._leer([("Router", "admin", "S3creta!", "192.168.1.1", "", "", "")])
        self.assertEqual(avisos, [])
        f = filas[0]
        self.assertEqual((f["fila"], f["resultado"], f["estado"]), (2, "nueva", "activa"))
        self.assertEqual(f["url"], "https://192.168.1.1")
        self.assertNotIn("S3creta!", f["secreto_cifrado"])
        self.assertEqual(_descifrar(f["secreto_cifrado"]), "S3creta!")

    def test_falta_titulo_o_contrasena_es_error(self):
        filas, _ = self._leer([("", "u", "x"), ("T", "u", "")])
        self.assertEqual([f["resultado"] for f in filas], ["error", "error"])
        self.assertEqual(filas[0]["detalle"], "Falta el título")
        self.assertEqual(filas[1]["detalle"], "Falta la contraseña")
        self.assertEqual(filas[1]["secreto_cifrado"], "")

    def test_ya_existe_en_la_boveda_por_titulo_y_usuario_sin_mayusculas(self):
        Credencial.objects.create(vault=self.vault, titulo="Router", usuario="Admin")
        filas, _ = self._leer([(" router ", "admin", "x")])
        self.assertEqual(filas[0]["resultado"], "ya_existe")
        self.assertEqual(filas[0]["secreto_cifrado"], "")
        from passwords.importacion import leer_credenciales
        otra, _ = leer_credenciales(_xlsx([(" router ", "admin", "x")]), _vault("Otra"))[:2]
        self.assertEqual(otra[0]["resultado"], "nueva")

    def test_repetida_dentro_del_archivo(self):
        filas, _ = self._leer([("Router", "admin", "a"), ("ROUTER", "Admin ", "b")])
        self.assertEqual([f["resultado"] for f in filas], ["nueva", "repetida"])
        self.assertEqual(filas[1]["detalle"], "Repetida en el archivo")

    def test_encabezados_alternativos(self):
        filas, avisos = self._leer([("NAS", "root", "x", "nas.local")], encabezados=("Nombre", "User", "Clave", "Sitio"))
        self.assertEqual(avisos, [])
        self.assertEqual((filas[0]["titulo"], filas[0]["usuario"], filas[0]["resultado"]), ("NAS", "root", "nueva"))

    def test_clave_numerica_no_pierde_formato(self):
        filas, _ = self._leer([("A", "u1", 123456), ("B", "u2", 123456.0)])
        self.assertEqual([_descifrar(f["secreto_cifrado"]) for f in filas], ["123456", "123456"])

    def test_clave_convertida_en_fecha_es_error(self):
        filas, _ = self._leer([("A", "u", date(2024, 2, 1))])
        self.assertEqual(filas[0]["resultado"], "error")
        self.assertIn("fecha", filas[0]["detalle"])

    def test_estado_y_fecha(self):
        filas, _ = self._leer([
            ("A", "u1", "x", "", "", "Revocada", "31/12/2027"),
            ("B", "u2", "x", "", "", "Bloqueada", ""),
            ("C", "u3", "x", "", "", "", "31-31-2027"),
            ("D", "u4", "x", "", "", "expirada", date(2027, 1, 15)),
        ])
        self.assertEqual((filas[0]["estado"], filas[0]["fecha_expiracion"]), ("revocada", "2027-12-31"))
        self.assertEqual(filas[1]["resultado"], "error")
        self.assertIn("Estado no válido", filas[1]["detalle"])
        self.assertEqual(filas[2]["resultado"], "error")
        self.assertIn("Fecha de expiración no válida", filas[2]["detalle"])
        self.assertEqual((filas[3]["estado"], filas[3]["fecha_expiracion"]), ("expirada", "2027-01-15"))

    def test_url_invalida_va_a_notas(self):
        filas, _ = self._leer([("A", "u", "x", "intranet de contabilidad", "Nota previa")])
        self.assertEqual(filas[0]["resultado"], "nueva")
        self.assertEqual(filas[0]["url"], "")
        self.assertEqual(filas[0]["notas"], "Nota previa\nSitio: intranet de contabilidad")

    def test_columna_contrasena_faltante_advierte(self):
        filas, avisos = self._leer([("A", "u")], encabezados=("Título", "Usuario"))
        self.assertEqual(avisos, ['No se encontró la columna "Contraseña" en la hoja «Sheet» (se lee solo la primera hoja).'])
        self.assertEqual(filas[0]["resultado"], "error")

    def test_filas_vacias_se_omiten_y_hay_limite(self):
        filas, _ = self._leer([("A", "u", "x"), (None, None, None), ("B", "u", "y")])
        self.assertEqual([f["fila"] for f in filas], [2, 4])
        from passwords.importacion import ArchivoInvalido
        with patch("passwords.importacion.MAX_FILAS", 1):
            with self.assertRaises(ArchivoInvalido):
                self._leer([("A", "u", "x"), ("B", "u", "y")])

    def test_contrasena_conserva_espacios(self):
        filas, _ = self._leer([("A", "u", " a b ")])
        self.assertEqual(_descifrar(filas[0]["secreto_cifrado"]), " a b ")

    def test_titulo_muy_largo_es_error(self):
        filas, _ = self._leer([("T" * 151, "u", "x")])
        self.assertEqual(filas[0]["detalle"], "El título supera 150 caracteres")

    def test_resumen(self):
        from passwords.importacion import resumen
        filas = [{"resultado": r} for r in ("nueva", "nueva", "ya_existe", "repetida", "error")]
        self.assertEqual(resumen(filas), {"nuevas": 2, "ya_existen": 2, "errores": 1, "total": 5})


# ─── Importación desde Excel: plantilla ──────────────────────────


class PlantillaImportacionTests(TestCase):
    def setUp(self):
        self.client.force_login(_user("superadmin"))

    def _descargar(self):
        from openpyxl import load_workbook
        resp = self.client.get(reverse("passwords:importar_plantilla"))
        self.assertEqual(resp.status_code, 200)
        self.assertIn("plantilla_importar_contrasenas.xlsx", resp["Content-Disposition"])
        return resp, load_workbook(BytesIO(resp.content))

    def test_hojas_encabezados_y_formato_texto(self):
        _, wb = self._descargar()
        self.assertEqual(wb.sheetnames, ["Plantilla Contraseñas", "Instrucciones"])
        ws = wb["Plantilla Contraseñas"]
        self.assertEqual([c.value for c in ws[1]], list(ENCABEZADOS))
        self.assertEqual(ws["C2"].number_format, "@")
        self.assertEqual(ws["C500"].number_format, "@")

    def test_filas_de_ejemplo_de_la_plantilla_no_se_importan(self):
        from openpyxl import load_workbook
        from passwords.importacion import leer_credenciales
        resp, _ = self._descargar()
        wb = load_workbook(BytesIO(resp.content))
        wb["Plantilla Contraseñas"].append(["Router oficina", "admin", "Clave-Real-1", "", "", "", ""])
        archivo = BytesIO()
        wb.save(archivo)
        archivo.seek(0)
        archivo.name = "plantilla.xlsx"
        filas, avisos = leer_credenciales(archivo, _vault())[:2]
        self.assertEqual([f["resultado"] for f in filas], ["error", "error", "nueva"])
        self.assertEqual(filas[0]["detalle"], "Fila de ejemplo de la plantilla: bórrala o reemplázala")
        self.assertEqual(_descifrar(filas[2]["secreto_cifrado"]), "Clave-Real-1")

    def test_admin_con_lectura_no_descarga(self):
        self.client.force_login(_user("admin"))
        resp = self.client.get(reverse("passwords:importar_plantilla"))
        self.assertRedirects(resp, reverse("core:dashboard"))


# ─── Importación desde Excel: vistas ─────────────────────────────


class ImportacionVistasTests(TestCase):
    CLAVE = "Clave-Secreta-XYZ"

    def setUp(self):
        self.client.force_login(_user("superadmin"))
        self.vault = _vault()

    def _url(self, nombre, vault=None):
        return reverse(f"passwords:{nombre}", kwargs={"vault_id": (vault or self.vault).pk})

    def _confirmar(self, vault=None, **kw):
        datos = self.client.session.get(f"importar_credenciales:{(vault or self.vault).pk}") or {}
        return self.client.post(self._url("credenciales_importar_confirmar", vault), {"lote": datos.get("lote", "")}, **kw)

    def _previsualizar(self, filas=None, vault=None, nombre="claves.xlsx"):
        filas = filas or [("Router", "admin", self.CLAVE, "192.168.1.1")]
        return self.client.post(self._url("credenciales_importar", vault), {"archivo": _xlsx(filas, nombre=nombre)})

    def test_preview_no_guarda_nada_ni_muestra_la_clave(self):
        resp = self._previsualizar()
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(Credencial.objects.count(), 0)
        self.assertNotContains(resp, self.CLAVE)
        self.assertContains(resp, "••••••••")
        self.assertNotIn(self.CLAVE, json.dumps(dict(self.client.session)))

    def test_confirmar_crea_credenciales_cifradas_y_registra_log(self):
        self._previsualizar()
        resp = self._confirmar()
        self.assertRedirects(resp, self._url("credenciales_lista"))
        cred = Credencial.objects.get(vault=self.vault)
        self.assertEqual(cred.get_secret(), self.CLAVE)
        self.assertEqual(cred.url, "https://192.168.1.1")
        log = AccesoLog.objects.get(vault=self.vault, accion="importacion_excel")
        self.assertEqual(log.detalle, "claves.xlsx: 1 creada, 0 omitidas")
        self.assertNotIn(self.CLAVE, log.detalle)

    def test_confirmar_sin_preview_avisa_que_expiro(self):
        resp = self.client.post(self._url("credenciales_importar_confirmar"), follow=True)
        self.assertContains(resp, "La vista previa expiró")
        self.assertEqual(Credencial.objects.count(), 0)

    def test_segunda_confirmacion_avisa_que_la_vista_previa_ya_no_existe(self):
        # La idempotencia por lote la cubren test_confirmar_con_sesion_vieja_avisa_que_ya_se_aplico
        # y ImportacionConcurrenteTests; aquí la sesión ya no tiene la vista previa.
        self._previsualizar()
        self._confirmar()
        resp = self._confirmar(follow=True)
        self.assertContains(resp, "La vista previa expiró")
        self.assertEqual(Credencial.objects.count(), 1)

    def test_no_sobrescribe_existentes(self):
        vieja = Credencial(vault=self.vault, titulo="Router", usuario="admin")
        vieja.set_secret("vieja")
        vieja.save()
        self._previsualizar()
        self._confirmar()
        self.assertEqual(Credencial.objects.count(), 1)
        vieja.refresh_from_db()
        self.assertEqual(vieja.get_secret(), "vieja")

    def test_confirmar_es_por_boveda(self):
        otra = _vault("Otra")
        self._previsualizar()
        resp = self.client.post(self._url("credenciales_importar_confirmar", otra), follow=True)
        self.assertContains(resp, "La vista previa expiró")
        self.assertEqual(Credencial.objects.count(), 0)

    def test_rechaza_archivo_que_no_es_xlsx(self):
        resp = self._previsualizar(nombre="claves.csv")
        self.assertRedirects(resp, self._url("credenciales_importar"))
        self.assertNotIn(f"importar_credenciales:{self.vault.pk}", self.client.session)

    def test_rechaza_archivo_muy_grande(self):
        with patch("passwords.importacion.MAX_BYTES", 10):
            resp = self._previsualizar()
        self.assertRedirects(resp, self._url("credenciales_importar"))
        self.assertNotIn(f"importar_credenciales:{self.vault.pk}", self.client.session)

    def test_admin_con_lectura_no_puede_importar(self):
        self.client.force_login(_user("admin"))
        self.assertRedirects(self.client.get(self._url("credenciales_importar")), reverse("core:dashboard"))

    def test_boton_importar_solo_con_escritura(self):
        url = self._url("credenciales_importar")
        self.assertContains(self.client.get(self._url("credenciales_lista")), url)
        self.client.force_login(_user("admin"))
        self.assertNotContains(self.client.get(self._url("credenciales_lista")), url)


# ─── Importación desde Excel: endurecimiento tras la revisión ────


class ImportacionEndurecimientoTests(TestCase):
    def setUp(self):
        self.vault = _vault()

    def _leer_wb(self, wb, nombre="claves.xlsx"):
        from passwords.importacion import leer_credenciales
        buffer = BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        buffer.name = nombre
        return leer_credenciales(buffer, self.vault)[:2]

    def _wb(self, filas, encabezados=ENCABEZADOS):
        wb = Workbook()
        ws = wb.active
        ws.append(list(encabezados))
        for fila in filas:
            ws.append(list(fila))
        return wb

    def test_celda_fantasma_no_agranda_la_lectura(self):
        # Una celda con formato en XFD1048576 (última fila y columna de Excel) declara
        # una hoja gigante. Se mide lo que openpyxl entrega, sin reloj: filas hasta la
        # última de Excel y nunca más columnas que las del encabezado.
        from passwords.importacion import MAX_COLUMNAS
        wb = self._wb([("A", "u", "x")])
        wb.active.cell(row=1048576, column=16384).number_format = "@"
        with _lectura_acotada(max_filas=1048576, max_ancho=MAX_COLUMNAS):
            filas, _ = self._leer_wb(wb)
        self.assertEqual([f["resultado"] for f in filas], ["nueva"])

    def test_valores_transformados_por_excel_son_error(self):
        wb = self._wb([
            ("Bool", "u1", True),
            ("Hora", "u2", time(12, 30)),
            ("Duracion", "u3", timedelta(days=1, hours=2)),
            ("Porcentaje", "u4", 0.5),
            ("Moneda", "u5", 100),
            ("Larga", "u6", 1234567890123456789),
            ("ErrorExcel", "u7", "#N/A"),
        ])
        ws = wb.active
        ws["C5"].number_format = "0%"
        ws["C6"].number_format = '"$"#,##0'
        filas, _ = self._leer_wb(wb)
        self.assertEqual([f["resultado"] for f in filas], ["error"] * 7)
        self.assertTrue(all(f["secreto_cifrado"] == "" for f in filas))
        self.assertIn("VERDADERO/FALSO", filas[0]["detalle"])
        self.assertIn("fecha u hora", filas[1]["detalle"])
        self.assertIn("fecha u hora", filas[2]["detalle"])
        self.assertIn("formato", filas[3]["detalle"])
        self.assertIn("formato", filas[4]["detalle"])
        self.assertIn("perdió dígitos", filas[5]["detalle"])
        self.assertIn("error de Excel", filas[6]["detalle"])

    def test_numero_valido_queda_nueva_con_aviso(self):
        filas, _ = self._leer_wb(self._wb([("A", "u", 123456)]))
        self.assertEqual(filas[0]["resultado"], "nueva")
        self.assertIn("número", filas[0]["detalle"])

    def test_encabezados_no_reconocidos_avisan(self):
        wb = self._wb([("Outlook", "a@x.com", "c1", "10.0.0.1", "dato")],
                      encabezados=("Servicio", "Correo electrónico", "Contraseña", "Dirección IP", "Extra"))
        filas, avisos = self._leer_wb(wb)
        self.assertEqual((filas[0]["usuario"], filas[0]["url"]), ("a@x.com", "https://10.0.0.1"))
        self.assertEqual(len(avisos), 1)
        self.assertIn("Extra", avisos[0])

    def test_alias_ampliados(self):
        wb = self._wb([("A", "u@x.com", "c", "intranet.local", "nota")],
                      encabezados=("Título", "Usuario/Correo", "Contraseñas", "Página web", "Observación"))
        filas, avisos = self._leer_wb(wb)
        self.assertEqual(avisos, [])
        f = filas[0]
        self.assertEqual((f["usuario"], f["url"], f["notas"], f["resultado"]),
                         ("u@x.com", "https://intranet.local", "nota", "nueva"))

    def test_sin_columna_usuario_avisa(self):
        filas, avisos = self._leer_wb(self._wb([("A", "c")], encabezados=("Título", "Contraseña")))
        self.assertEqual(avisos, ['No se encontró la columna "Usuario": las filas con el mismo título '
                                  'se tratarán como repetidas.'])

    def test_columnas_en_otro_orden_no_importa_nada(self):
        wb = self._wb([
            ("NAS", "admin", "https://nas.local", "Nas.Backup#7"),
            ("FW", "root", "192.168.1.1", "Fw-Clave-2024!"),
            ("Router", "admin", "http://192.168.1.254", "Rt-Clave-9"),
        ])
        filas, avisos = self._leer_wb(wb)
        self.assertTrue(any("orden" in a for a in avisos))
        self.assertEqual({f["resultado"] for f in filas}, {"error"})
        self.assertEqual({f["detalle"] for f in filas}, {"Revisa el orden de las columnas"})
        self.assertTrue(all(f["secreto_cifrado"] == "" and f["url"] == "" and f["notas"] == "" for f in filas))

    def test_url_que_no_es_direccion_se_marca_en_la_fila(self):
        filas, _ = self._leer_wb(self._wb([("A", "u", "x", "intranet de contabilidad")]))
        self.assertTrue(filas[0]["url_a_notas"])


class ImportacionVistasEndurecimientoTests(TestCase):
    def setUp(self):
        self.client.force_login(_user("superadmin"))
        self.vault = _vault()
        self.clave_sesion = f"importar_credenciales:{self.vault.pk}"

    def _url(self, nombre):
        return reverse(f"passwords:{nombre}", kwargs={"vault_id": self.vault.pk})

    def _previsualizar(self, filas):
        return self.client.post(self._url("credenciales_importar"), {"archivo": _xlsx(filas)})

    def test_confirmar_con_sesion_vieja_avisa_que_ya_se_aplico(self):
        self._previsualizar([("Router", "admin", "c1")])
        datos = dict(self.client.session[self.clave_sesion])
        self.client.post(self._url("credenciales_importar_confirmar"), {"lote": datos["lote"]})
        sesion = self.client.session
        sesion[self.clave_sesion] = datos
        sesion.save()
        resp = self.client.post(self._url("credenciales_importar_confirmar"), {"lote": datos["lote"]}, follow=True)
        self.assertContains(resp, "Esta importación ya se había aplicado")
        self.assertEqual(Credencial.objects.count(), 1)
        self.assertEqual(AccesoLog.objects.filter(accion="importacion_excel").count(), 1)

    def test_lote_distinto_no_importa(self):
        self._previsualizar([("Router", "admin", "c1")])
        resp = self.client.post(self._url("credenciales_importar_confirmar"), {"lote": "otro"}, follow=True)
        self.assertContains(resp, "La vista previa cambió")
        self.assertEqual(Credencial.objects.count(), 0)

    def test_sesion_guarda_solo_las_nuevas(self):
        self._previsualizar([("A", "u", "c1"), ("", "u", "c2"), ("A", "u", "c3")])
        datos = self.client.session[self.clave_sesion]
        self.assertEqual(sorted(datos), ["archivo", "lote", "nuevas", "total"])
        self.assertEqual((len(datos["nuevas"]), datos["total"]), (1, 3))

    def test_preview_marca_url_a_notas_sin_mostrarla(self):
        resp = self._previsualizar([("A", "u", "x", "intranet de contabilidad")])
        self.assertContains(resp, "→ Notas")
        self.assertNotContains(resp, "intranet de contabilidad")

    def test_error_de_lectura_no_deja_la_celda_en_los_logs(self):
        error = ValueError("could not convert string to float: 'Clave-Real-9!'")
        with patch("passwords.importacion.leer_credenciales", side_effect=error):
            with self.assertLogs("passwords.views", level="WARNING") as registro:
                resp = self._previsualizar([("A", "u", "x")])
        self.assertRedirects(resp, self._url("credenciales_importar"))
        self.assertNotIn("Clave-Real-9!", "\n".join(registro.output))


# ─── Importación desde Excel: arreglos de la auditoría final ─────


class ImportacionColumnasCruzadasTests(TestCase):
    def setUp(self):
        self.vault = _vault()

    def _leer(self, filas, encabezados=ENCABEZADOS):
        from passwords.importacion import leer_credenciales
        return leer_credenciales(_xlsx(filas, encabezados=encabezados), self.vault)[:2]

    def test_servidores_sin_dominio_no_bloquean_el_archivo(self):
        filas, avisos = self._leer([
            ("Controlador de dominio", "administrador", "Dc-Clave-2024!", "SRV-DC01"),
            ("Archivos", "administrador", "Fs-Clave-2024!", "SRV-FS01"),
            ("Base de datos", "sa", "Db-Clave-2024!", "SRV-DB01"),
            ("Impresora", "admin", "Imp-Clave-1", "192.168.10.20"),
        ], encabezados=("Título", "Usuario", "Contraseña", "Servidor"))
        self.assertEqual(avisos, [])
        self.assertEqual([f["resultado"] for f in filas], ["nueva"] * 4)
        self.assertEqual(filas[0]["notas"], "Sitio: SRV-DC01")
        self.assertEqual(filas[3]["url"], "https://192.168.10.20")

    def test_una_sola_clave_con_forma_de_direccion_no_bloquea(self):
        filas, avisos = self._leer([("Wifi casa", "", "www.Casa2024")])
        self.assertEqual(avisos, [])
        self.assertEqual(filas[0]["resultado"], "nueva")
        self.assertEqual(_descifrar(filas[0]["secreto_cifrado"]), "www.Casa2024")

    def test_columna_sitio_con_descripciones_no_bloquea(self):
        filas, avisos = self._leer([
            ("Caja fuerte", "u1", "Caja-1", "Oficina principal"),
            ("Alarma", "u2", "Alarma-2", "Bodega norte"),
            ("Portón", "u3", "Porton-3", "Sede centro"),
        ], encabezados=("Título", "Usuario", "Contraseña", "Sitio"))
        self.assertEqual(avisos, [])
        self.assertEqual([f["resultado"] for f in filas], ["nueva"] * 3)
        self.assertEqual(filas[1]["notas"], "Sitio: Bodega norte")

    def test_el_aviso_dice_que_columna_y_por_que(self):
        filas, avisos = self._leer([
            ("NAS", "admin", "https://nas.local", "Nas.Backup#7"),
            ("FW", "root", "192.168.1.1", "Fw-Clave-2024!"),
            ("Router", "admin", "www.router.local", "Rt-Clave-9"),
        ], encabezados=("Título", "Usuario", "Clave", "URL"))
        aviso = next(a for a in avisos if "«Clave»" in a)
        self.assertIn("3 de 3", aviso)
        self.assertIn("parecen direcciones", aviso)
        self.assertEqual({f["resultado"] for f in filas}, {"error"})

    def test_urls_que_parecen_claves_bloquean_y_nombran_la_columna(self):
        filas, avisos = self._leer([
            ("NAS", "admin", "Clave-Uno-1", "Nas.Backup#7"),
            ("FW", "root", "Clave-Dos-2", "Fw$Clave2024!"),
            ("Router", "admin", "Clave-Tres-3", "Rt*Clave9"),
        ], encabezados=("Título", "Usuario", "Contraseña", "Servidor"))
        aviso = next(a for a in avisos if "«Servidor»" in a)
        self.assertIn("3 de 3", aviso)
        self.assertEqual({f["resultado"] for f in filas}, {"error"})
        self.assertTrue(all(f["secreto_cifrado"] == "" and f["notas"] == "" for f in filas))

    def test_fila_corrida_con_clave_de_direccion_no_se_importa(self):
        # Una fila desplazada entre filas correctas: la clave real quedó en URL y, con
        # forma de host, se guardaría como URL visible para el rol lectura.
        filas = [(f"Servicio {i}", "u", f"Clave-{i}-Ok", "https://ok.example.com") for i in range(10)]
        filas.append(("Firewall", "admin", "https://fw.example.com", "Clave.Segura"))
        leidas, avisos = self._leer(filas)
        self.assertEqual(avisos, [])
        self.assertEqual([f["resultado"] for f in leidas], ["nueva"] * 10 + ["error"])
        corrida = leidas[-1]
        self.assertIn("columnas corridas", corrida["detalle"])
        self.assertEqual((corrida["url"], corrida["notas"], corrida["secreto_cifrado"]), ("", "", ""))

    def test_fila_corrida_con_url_que_no_es_direccion_no_se_importa(self):
        filas = [(f"Servicio {i}", "u", f"Clave-{i}-Ok", "https://ok.example.com") for i in range(10)]
        filas.append(("Firewall", "admin", "https://fw.example.com", "Nas-Backup#7"))
        leidas, _ = self._leer(filas)
        self.assertEqual(leidas[-1]["resultado"], "error")
        self.assertEqual(leidas[-1]["notas"], "")

    def test_url_con_terminacion_desconocida_pide_revisar(self):
        filas, _ = self._leer([("Firewall", "admin", "Fw-Clave-1", "Clave.Segura")])
        self.assertEqual(filas[0]["resultado"], "nueva")
        self.assertIn("Revisa", filas[0]["detalle"])
        self.assertIn("URL", filas[0]["detalle"])

    def test_valor_de_url_que_queda_como_nota_pide_revisar(self):
        # Residuo de la auditoría: con Contraseña y Servidor intercambiadas, una clave
        # sin símbolos (Casa2024) tiene la misma forma que un nombre de equipo, así que
        # no bloquea el archivo y la clave real queda en notas, un TextField sin cifrar.
        # No se puede bloquear (volvería el falso positivo I1), pero la fila avisa.
        filas, avisos = self._leer([
            ("Wifi casa", "", "casa.example.com", "Casa2024"),
        ], encabezados=("Título", "Usuario", "Contraseña", "Servidor"))
        self.assertEqual(avisos, [])
        self.assertEqual(filas[0]["resultado"], "nueva")
        self.assertEqual(filas[0]["notas"], "Sitio: Casa2024")
        self.assertIn("Revisa", filas[0]["detalle"])
        self.assertIn("«Servidor»", filas[0]["detalle"])
        self.assertIn("nota", filas[0]["detalle"].lower())
        self.assertNotIn("Casa2024", filas[0]["detalle"])

    def test_descripcion_con_espacios_que_queda_como_nota_no_pide_revisar(self):
        # "Oficina principal" no puede ser una contraseña corrida de columna: lleva
        # espacios. Queda como nota, sin aviso, para no llenar de avisos un archivo bueno.
        filas, _ = self._leer([
            ("Caja fuerte", "u1", "Caja-1", "Oficina principal"),
        ], encabezados=("Título", "Usuario", "Contraseña", "Sitio"))
        self.assertEqual(filas[0]["resultado"], "nueva")
        self.assertEqual(filas[0]["notas"], "Sitio: Oficina principal")
        self.assertEqual(filas[0]["detalle"], "")

    def test_urls_normales_no_piden_revisar(self):
        # SRV-DC01 salió de esta lista: no es una dirección válida, queda como nota sin
        # cifrar y por eso ahora pide revisar (ver el test del residuo, arriba).
        urls = ["nas.local", "fw.example.com", "outlook.office.com", "192.168.1.1", "erp.empresa.co",
                "https://intranet.empresa", "localhost:8080"]
        filas, avisos = self._leer([(f"T{i}", "u", f"Clave-{i}", url) for i, url in enumerate(urls)])
        self.assertEqual(avisos, [])
        self.assertEqual([f["detalle"] for f in filas], [""] * len(urls))


class ImportacionArchivoFabricadoTests(TestCase):
    def setUp(self):
        self.vault = _vault()

    def _leer(self, archivo):
        from passwords.importacion import leer_credenciales
        return leer_credenciales(archivo, self.vault)[:2]

    def test_fila_mas_alla_de_excel_se_rechaza_sin_recorrer_el_hueco(self):
        from passwords.importacion import ArchivoInvalido
        archivo = _xlsx_a_mano('<row r="2000000000"><c r="A2000000000" t="inlineStr"><is><t>x</t></is></c></row>')
        self.assertLess(len(archivo.getvalue()), 2048)
        # 1 fila de encabezado + las filas 2 a 1.048.577: ni una más.
        with _lectura_acotada(max_filas=1048577, max_ancho=50):
            with self.assertRaisesMessage(ArchivoInvalido, "1.048.576"):
                self._leer(archivo)

    def test_shared_strings_inflado_se_rechaza_antes_de_cargarlo(self):
        from openpyxl.reader.excel import ExcelReader
        from passwords.importacion import MAX_BYTES, ArchivoInvalido
        archivo = _xlsx_a_mano(shared_strings="<si><t>" + "A" * MAX_BYTES + "</t></si>")
        self.assertLess(len(archivo.getvalue()), 64 * 1024)
        with patch.object(ExcelReader, "read_strings", side_effect=AssertionError("openpyxl cargó sharedStrings")):
            with self.assertRaises(ArchivoInvalido):
                self._leer(archivo)

    def test_archivo_que_se_infla_al_descomprimir_se_rechaza(self):
        from passwords.importacion import MAX_BYTES, ArchivoInvalido
        relleno = {f"docProps/relleno{i}.bin": b"\0" * MAX_BYTES for i in range(5)}
        with self.assertRaises(ArchivoInvalido):
            self._leer(_xlsx_a_mano(partes_extra=relleno))

    def test_solo_se_prepara_la_primera_hoja(self):
        # Al abrir, openpyxl recorre cada hoja declarada: 200 hojas que apuntan a la
        # misma parte multiplican por 200 el trabajo con un archivo de pocos KB.
        from openpyxl.worksheet._read_only import ReadOnlyWorksheet
        archivo = _xlsx_a_mano(
            '<row r="2"><c r="A2" t="inlineStr"><is><t>Router</t></is></c>'
            '<c r="B2" t="inlineStr"><is><t>Clave-1</t></is></c></row>',
            hojas_declaradas=200,
        )
        with patch.object(ReadOnlyWorksheet, "_get_size", autospec=True,
                          side_effect=ReadOnlyWorksheet._get_size) as preparar:
            filas, _ = self._leer(archivo)
        self.assertEqual(preparar.call_count, 1)
        self.assertEqual([f["resultado"] for f in filas], ["nueva"])


class ImportacionValoresTests(TestCase):
    def setUp(self):
        self.vault = _vault()

    def _leer(self, filas, **kw):
        from passwords.importacion import leer_credenciales
        return leer_credenciales(_xlsx(filas, **kw), self.vault)

    def test_decimales_que_excel_no_muestra_igual_son_error(self):
        # Excel guarda con 17 cifras el resultado de =0,1+0,2 (y muestra 0,3) y en
        # notación científica los números muy chicos; openpyxl no escribe así.
        from passwords.importacion import leer_credenciales
        filas_xml = "".join(
            f'<row r="{n}"><c r="A{n}" t="inlineStr"><is><t>T{n}</t></is></c><c r="B{n}"><v>{v}</v></c></row>'
            for n, v in ((2, "0.30000000000000004"), (3, "9.9999999999999995E-8"), (4, "3.25"))
        )
        filas, _ = leer_credenciales(_xlsx_a_mano(filas_xml), self.vault)[:2]
        self.assertEqual([f["resultado"] for f in filas], ["error", "error", "nueva"])
        self.assertIn("decimales", filas[0]["detalle"])
        self.assertEqual(filas[0]["secreto_cifrado"], "")
        self.assertEqual(_descifrar(filas[2]["secreto_cifrado"]), "3.25")

    def test_fecha_con_solo_espacios_queda_vacia(self):
        filas, _ = self._leer([("A", "u", "x", "", "", "", "   ")])[:2]
        self.assertEqual((filas[0]["resultado"], filas[0]["fecha_expiracion"]), ("nueva", None))

    def test_errores_de_estado_y_fecha_no_repiten_el_valor(self):
        filas, _ = self._leer([
            ("Router", "admin", "Activa", "", "", "Clave-En-Estado-55"),
            ("Switch", "admin", "Activa", "", "", "", "Clave-En-Fecha-44"),
        ])[:2]
        self.assertIn("Estado no válido", filas[0]["detalle"])
        self.assertIn("Fecha de expiración no válida", filas[1]["detalle"])
        self.assertNotIn("clave-en-estado-55", filas[0]["detalle"].lower())
        self.assertNotIn("clave-en-fecha-44", filas[1]["detalle"].lower())

    def test_archivo_sin_encabezados_no_lista_los_valores_de_la_fila_1(self):
        _, avisos = self._leer(
            [("Router", "admin", "Clave-Fila-2")],
            encabezados=("Correo Gerencia", "gerencia@empresa.com", "Clave-Primera-Fila-99"),
        )[:2]
        texto = " ".join(avisos)
        self.assertNotIn("Clave-Primera-Fila-99", texto)
        self.assertNotIn("gerencia@empresa.com", texto)
        self.assertIn("A, B y C", texto)

    def test_dice_que_hoja_se_leyo(self):
        wb = Workbook()
        wb.active.title = "Instrucciones"
        wb.active.append(["Lea esto primero"])
        datos = wb.create_sheet("Contraseñas")
        datos.append(list(ENCABEZADOS))
        datos.append(["Router", "admin", "Clave-1"])
        buffer = BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        from passwords.importacion import leer_credenciales
        _, avisos, hoja = leer_credenciales(buffer, self.vault)
        self.assertEqual(hoja, "Instrucciones")
        self.assertIn('No se encontró la columna "Título" en la hoja «Instrucciones»', avisos[0])


class ImportacionVistasAuditoriaTests(TestCase):
    def setUp(self):
        self.usuario = _user("superadmin")
        self.client.force_login(self.usuario)
        self.vault = _vault()
        self.clave_sesion = f"importar_credenciales:{self.vault.pk}"

    def _url(self, nombre):
        return reverse(f"passwords:{nombre}", kwargs={"vault_id": self.vault.pk})

    def _previsualizar(self, archivo, cliente=None, **extra):
        return (cliente or self.client).post(
            self._url("credenciales_importar"), {"archivo": archivo, **extra}, follow=True,
        )

    def _cliente_con_csrf(self):
        cliente = Client(enforce_csrf_checks=True)
        cliente.force_login(self.usuario)
        token = str(cliente.get(self._url("credenciales_importar")).context["csrf_token"])
        return cliente, token

    def test_preview_no_muestra_valores_que_pueden_ser_claves(self):
        resp = self._previsualizar(_xlsx([
            ("Router", "admin", "Activa", "", "", "Clave-En-Estado-55"),
            ("Switch", "admin", "Activa", "", "", "", "Clave-En-Fecha-44"),
        ]))
        self.assertNotContains(resp, "Clave-En-Estado-55")
        self.assertNotContains(resp, "Clave-En-Fecha-44")
        resp = self._previsualizar(_xlsx(
            [("Router", "admin", "Clave-Fila-2")],
            encabezados=("Correo Gerencia", "gerencia@empresa.com", "Clave-Primera-Fila-99"),
        ))
        self.assertNotContains(resp, "Clave-Primera-Fila-99")

    def test_vista_previa_que_no_cabe_en_la_sesion_se_rechaza(self):
        notas = "N" * 25000
        resp = self._previsualizar(_xlsx([(f"T{i}", "u", f"Clave-{i}", "", notas) for i in range(100)]))
        self.assertContains(resp, "ocupan más de 2 MB")
        self.assertNotIn(self.clave_sesion, self.client.session)

    def test_mil_filas_normales_caben_en_la_sesion(self):
        nota = "Cuenta compartida del área; pedir acceso a Sistemas. " * 4
        filas = [(f"Servicio {i}", f"usuario{i}@empresa.com", f"Clave-{i:04d}-Segura!",
                  f"https://app{i}.empresa.com", nota, "Activa", "31/12/2027") for i in range(1000)]
        resp = self._previsualizar(_xlsx(filas))
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(len(self.client.session[self.clave_sesion]["nuevas"]), 1000)

    def test_subida_de_mas_de_6_mb_se_rechaza_sin_archivo_temporal(self):
        cliente, token = self._cliente_con_csrf()
        grande = BytesIO(b"\0" * (settings.FILE_UPLOAD_MAX_MEMORY_SIZE + 1024))
        grande.name = "grande.xlsx"
        with patch("django.core.files.uploadhandler.TemporaryUploadedFile",
                   side_effect=AssertionError("la subida se escribió en un archivo temporal")):
            resp = self._previsualizar(grande, cliente, csrfmiddlewaretoken=token)
        self.assertContains(resp, "El archivo pasa de 5 MB")
        self.assertNotIn(self.clave_sesion, cliente.session)

    def test_el_excel_llega_a_la_vista_en_memoria(self):
        from django.core.files.uploadedfile import InMemoryUploadedFile
        cliente, token = self._cliente_con_csrf()
        recibido = {}

        def leer(archivo, vault):
            recibido["archivo"] = archivo
            return [], [], "Hoja1"

        with patch("passwords.importacion.leer_credenciales", side_effect=leer), \
                patch("django.core.files.uploadhandler.TemporaryUploadedFile",
                      side_effect=AssertionError("la subida se escribió en un archivo temporal")):
            resp = self._previsualizar(_xlsx([("A", "u", "x")]), cliente, csrfmiddlewaretoken=token)
        self.assertEqual(resp.status_code, 200)
        self.assertIsInstance(recibido["archivo"], InMemoryUploadedFile)
        self.assertFalse(hasattr(recibido["archivo"], "temporary_file_path"))

    def test_importar_sigue_exigiendo_csrf(self):
        cliente = Client(enforce_csrf_checks=True)
        cliente.force_login(self.usuario)
        resp = cliente.post(self._url("credenciales_importar"), {"archivo": _xlsx([("A", "u", "x")])})
        self.assertEqual(resp.status_code, 403)
        self.assertNotIn(self.clave_sesion, cliente.session)

    def test_confirmar_guarda_todos_los_campos(self):
        self._previsualizar(_xlsx([("NAS", "root", "Clave-Nas-1", "intranet de contabilidad",
                                    "Respaldo semanal", "Revocada", "31/12/2027")]))
        lote = self.client.session[self.clave_sesion]["lote"]
        self.client.post(self._url("credenciales_importar_confirmar"), {"lote": lote})
        cred = Credencial.objects.get(vault=self.vault)
        self.assertEqual(
            (cred.titulo, cred.usuario, cred.url, cred.notas, cred.estado, cred.fecha_expiracion, cred.creado_por),
            ("NAS", "root", "", "Respaldo semanal\nSitio: intranet de contabilidad", "revocada",
             date(2027, 12, 31), self.usuario),
        )
        self.assertEqual(cred.get_secret(), "Clave-Nas-1")

    def test_archivo_con_demasiadas_filas_avisa_por_la_vista(self):
        with patch("passwords.importacion.MAX_FILAS", 2):
            resp = self._previsualizar(_xlsx([("A", "u", "1"), ("B", "u", "2"), ("C", "u", "3")]))
        self.assertContains(resp, "más de 2 filas")
        self.assertNotIn(self.clave_sesion, self.client.session)

    def test_archivo_danado_avisa_sin_mostrar_su_contenido(self):
        danado = BytesIO(b"PK\x03\x04" + b"Clave-Dentro-Del-Archivo" * 10)
        danado.name = "danado.xlsx"
        with self.assertLogs("passwords.views", level="WARNING") as registro:
            resp = self._previsualizar(danado)
        self.assertContains(resp, "¿Es un .xlsx válido?")
        self.assertNotContains(resp, "Clave-Dentro-Del-Archivo")
        self.assertNotIn("Clave-Dentro-Del-Archivo", "\n".join(registro.output))

    def test_admin_con_lectura_no_puede_subir_ni_confirmar(self):
        self._previsualizar(_xlsx([("Router", "admin", "Clave-1")]))
        datos = dict(self.client.session[self.clave_sesion])
        admin = Client()
        admin.force_login(_user("admin"))
        sesion = admin.session
        sesion[self.clave_sesion] = datos
        sesion.save()
        resp = admin.post(self._url("credenciales_importar_confirmar"), {"lote": datos["lote"]})
        self.assertRedirects(resp, reverse("core:dashboard"))
        resp = admin.post(self._url("credenciales_importar"), {"archivo": _xlsx([("Switch", "admin", "Clave-2")])})
        self.assertRedirects(resp, reverse("core:dashboard"))
        self.assertEqual(Credencial.objects.count(), 0)

    def test_textos_en_singular_y_plural(self):
        resp = self._previsualizar(_xlsx([("Router", "admin", "Clave-1")]))
        self.assertContains(resp, "1 fila leída")
        self.assertContains(resp, "Importar 1 contraseña\n")
        self.assertNotContains(resp, "1 contraseñas")
        lote = self.client.session[self.clave_sesion]["lote"]
        resp = self.client.post(self._url("credenciales_importar_confirmar"), {"lote": lote}, follow=True)
        self.assertContains(resp, "Importación terminada: 1 contraseña creada, 0 omitidas.")
        resp = self._previsualizar(_xlsx([("A", "u", "c1"), ("B", "u", "c2")]))
        self.assertContains(resp, "2 filas leídas")
        self.assertContains(resp, "Importar 2 contraseñas")

    def test_preview_dice_que_hoja_se_leyo(self):
        wb = Workbook()
        wb.active.title = "Claves oficina"
        wb.active.append(list(ENCABEZADOS))
        wb.active.append(["Router", "admin", "Clave-1"])
        buffer = BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        buffer.name = "claves.xlsx"
        resp = self._previsualizar(buffer)
        self.assertContains(resp, "hoja «Claves oficina»")


@skipUnless(connection.vendor == "postgresql", "select_for_update solo bloquea en PostgreSQL")
class ImportacionConcurrenteTests(TransactionTestCase):
    def test_dos_confirmaciones_simultaneas_no_duplican(self):
        usuario = _user("superadmin")
        vault = _vault()

        def url(nombre):
            return reverse(f"passwords:{nombre}", kwargs={"vault_id": vault.pk})

        primero = Client()
        primero.force_login(usuario)
        primero.post(url("credenciales_importar"), {"archivo": _xlsx([(f"T{i}", "u", f"c{i}") for i in range(200)])})
        lote = primero.session[f"importar_credenciales:{vault.pk}"]["lote"]
        segundo = Client()
        segundo.cookies = primero.cookies
        barrera = threading.Barrier(2)

        def confirmar(cliente):
            try:
                barrera.wait()
                cliente.post(url("credenciales_importar_confirmar"), {"lote": lote})
            finally:
                connection.close()

        hilos = [threading.Thread(target=confirmar, args=(c,)) for c in (primero, segundo)]
        for hilo in hilos:
            hilo.start()
        for hilo in hilos:
            hilo.join()
        self.assertEqual(Credencial.objects.filter(vault=vault).count(), 200)
