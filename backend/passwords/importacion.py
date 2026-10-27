"""Importación masiva de credenciales desde Excel.

Cada contraseña se cifra en cuanto se lee: la vista previa y la sesión (que
vive en la base de datos) solo guardan el texto cifrado, nunca la clave.
"""
import json
import re
import unicodedata
import warnings
import zipfile
from datetime import date, datetime, time, timedelta
from urllib.parse import urlsplit

import openpyxl
from django.core.exceptions import ValidationError
from django.core.validators import URLValidator
from openpyxl.reader.excel import ExcelReader
from openpyxl.utils import get_column_letter

from .crypto import build_fernet
from .models import Credencial

MAX_FILAS = 1000
MAX_BYTES = 5 * 1024 * 1024
# Solo se miran las primeras columnas del encabezado: una celda con formato en
# XFD1 haría que el rango "usado" de la hoja tenga 16.384 columnas.
MAX_COLUMNAS = 50
# Última fila que existe en Excel. openpyxl rellena con filas vacías los huecos
# entre números de fila; más allá de esta solo hay archivos fabricados.
MAX_FILA_EXCEL = 1_048_576
# Un .xlsx es un zip: openpyxl carga enteras partes como sharedStrings.xml y
# styles.xml. 1.000 filas reales son unos cientos de KB de XML; estos topes dejan
# margen para hojas anchas u otras hojas y acotan el trabajo a unos segundos.
MAX_BYTES_PARTE = MAX_BYTES
MAX_BYTES_DESCOMPRIMIDOS = 4 * MAX_BYTES
# Lo que la vista previa guarda en la sesión (en la BD, reescrita en cada petición):
# 2 KB por fila en promedio con el máximo de filas. 1.000 filas normales usan ~0,6 MB.
MAX_BYTES_SESION = MAX_FILAS * 2 * 1024

ALIAS = {
    "titulo": {"titulo", "nombre", "servicio", "sistema", "aplicacion", "plataforma"},
    "usuario": {
        "usuario", "user", "username", "login", "cuenta", "correo", "email", "e mail", "mail",
        "correo electronico", "usuario correo", "nombre de usuario",
    },
    "secreto": {
        "contrasena", "contrasenas", "clave", "claves", "password", "passwords", "pass", "clave de acceso",
    },
    "url": {
        "url", "sitio", "sitio web", "enlace", "link", "direccion", "ip", "direccion ip",
        "pagina", "pagina web", "web", "host", "servidor",
    },
    "notas": {
        "notas", "nota", "observaciones", "observacion", "comentarios", "comentario", "descripcion", "detalle",
    },
    "estado": {"estado"},
    "fecha_expiracion": {
        "fecha de expiracion", "expiracion", "vence", "vencimiento", "fecha de vencimiento",
        "caduca", "caducidad",
    },
}
ESTADOS = {"activa", "expirada", "revocada"}
FORMATOS_FECHA = ("%d/%m/%Y", "%Y-%m-%d")
COLUMNAS_OBLIGATORIAS = (("titulo", "Título"), ("secreto", "Contraseña"))
# Formatos de número con los que Excel guarda lo mismo que se escribió.
FORMATOS_NUMERO_SEGUROS = {"General", "0", "@"}
# Un número de Excel (double) pierde dígitos a partir de ~15 cifras.
LIMITE_DIGITOS = 10 ** 15
# Solo se antepone https:// a algo con forma de host[:puerto][/ruta].
FORMA_DE_HOST = re.compile(r"[\w.-]+(:\d+)?(/\S*)?")
# Nombre de equipo sin dominio (SRV-DC01, nas:8080): no es una URL válida y va a
# Notas, pero en una columna "Servidor" o "Host" es lo normal, no una clave.
HOST_SIN_PUNTO = re.compile(r"[\w-]+(:\d+)?(/\S*)?")
PARECE_DIRECCION = re.compile(r"^(https?://|www\.|\\\\)|^\d{1,3}(\.\d{1,3}){3}(:\d+)?(/\S*)?$", re.IGNORECASE)
# Terminaciones con las que una URL escrita sin https:// se da por buena (además de
# las de dos letras: .co, .es…). Con otra ("Clave.Segura") la fila pide revisar.
DOMINIOS_CONOCIDOS = {
    "com", "net", "org", "edu", "gov", "mil", "int", "info", "biz", "io", "app", "dev", "cloud",
    "online", "site", "web", "local", "lan", "corp", "internal", "intranet", "home", "localdomain",
}
# Una columna entera solo se da por cruzada con al menos estas filas.
MIN_FILAS_ORDEN = 3

ENCABEZADOS_PLANTILLA = ["Título", "Usuario", "Contraseña", "URL", "Notas", "Estado", "Fecha de expiración"]
EJEMPLOS_PLANTILLA = [
    ["Correo Gerencia", "gerencia@empresa.com", "Ejemplo-Clave-1", "https://outlook.office.com",
     "Cuenta principal", "Activa", ""],
    ["Router oficina", "admin", "Ejemplo-Clave-2", "http://192.168.1.1", "", "Activa", "31/12/2027"],
]
CLAVES_EJEMPLO = {ejemplo[2] for ejemplo in EJEMPLOS_PLANTILLA}
INSTRUCCIONES_PLANTILLA = [
    "INSTRUCCIONES DE USO",
    "",
    "1. Pegue sus contraseñas en la hoja 'Plantilla Contraseñas', debajo de los encabezados.",
    "   Borre o reemplace las dos filas de ejemplo: si las deja, se marcan con error y no se importan.",
    "2. Todas las columnas están en formato Texto para que Excel no cambie claves como 00123 o 1/2/24.",
    "   Si pega desde otro Excel, use Pegado especial > Texto (o 'Mantener solo texto') para que se copie",
    "   lo que se ve. Evite 'Pegar valores' si alguna clave parece fecha, hora o porcentaje.",
    "3. Guarde como .xlsx y súbalo desde la bóveda: Importar desde Excel.",
    "",
    "COLUMNAS OBLIGATORIAS: Título y Contraseña.",
    "COLUMNAS OPCIONALES: Usuario, URL, Notas, Estado, Fecha de expiración.",
    "  - También se reconocen: Nombre o Servicio (Título); Correo, E-mail o Login (Usuario);",
    "    Clave o Password (Contraseña); Sitio, Página web o Dirección IP (URL); Observaciones (Notas).",
    "  - Estado: Activa, Expirada o Revocada (vacío = Activa).",
    "  - Fecha de expiración: dd/mm/aaaa (ej. 31/12/2027) o aaaa-mm-dd.",
    "  - URL: si no empieza por http(s):// se le agrega https://; si no es una dirección válida, va a Notas.",
    "",
    "IMPORTANTE:",
    "  - Si en la bóveda ya existe una credencial con el mismo Título y Usuario, la fila se omite (no se sobrescribe).",
    f"  - Máximo {MAX_FILAS} filas por archivo.",
    "  - Nada se guarda hasta que confirme en la vista previa.",
    "  - Cuando termine, borre este archivo de su PC y de la papelera: tiene las contraseñas en texto plano.",
]


class ArchivoInvalido(Exception):
    """El archivo no se puede importar; el texto es el mensaje para el usuario."""


def _sin_tildes(texto):
    descompuesto = unicodedata.normalize("NFKD", texto)
    return "".join(c for c in descompuesto if not unicodedata.combining(c))


def _normalizar_encabezado(valor):
    texto = _sin_tildes(str(valor or "")).lower()
    return " ".join(re.sub(r"[^a-z0-9]+", " ", texto).split())


def _campo_de(encabezado):
    normalizado = _normalizar_encabezado(encabezado)
    for campo, alias in ALIAS.items():
        if normalizado in alias:
            return campo
    return None


def _texto(valor):
    """Valor de celda como texto. Las fechas se tratan aparte."""
    if valor is None:
        return ""
    if isinstance(valor, bool):
        return str(valor)
    if isinstance(valor, int):
        return str(valor)
    if isinstance(valor, float):
        return str(int(valor)) if valor.is_integer() else str(valor)
    return str(valor)


def _error_de_la_clave(valor, tipo, formato):
    """Motivo por el que Excel pudo alterar la contraseña, o ""."""
    if isinstance(valor, bool) or tipo == "b":
        return "Excel convirtió la contraseña en VERDADERO/FALSO; escríbela como texto"
    if tipo == "e":
        return "La celda de la contraseña tiene un error de Excel (#N/A, #REF!…)"
    if isinstance(valor, (datetime, date, time, timedelta)):
        return "Excel convirtió la contraseña en fecha u hora; usa la plantilla (columnas en Texto)"
    if isinstance(valor, (int, float)):
        if abs(valor) >= LIMITE_DIGITOS:
            return "Número demasiado largo: Excel perdió dígitos; escribe la contraseña como texto"
        if (formato or "General") not in FORMATOS_NUMERO_SEGUROS:
            return ("Excel le dio formato de número (porcentaje, moneda o fecha) a la contraseña; "
                    "escríbela como texto")
        if isinstance(valor, float) and not valor.is_integer() and not _decimal_legible(valor):
            return ("Excel guardó la contraseña como un número con decimales que no muestra igual "
                    "(p. ej. 0,30000000000000004 o 1E-07); escríbela como texto")
    return ""


def _decimal_legible(valor):
    """True si el texto del número es el que Excel muestra: hasta 15 cifras y sin exponente."""
    texto = repr(valor)
    cifras = texto.lstrip("-").replace(".", "").lstrip("0")
    return "e" not in texto.lower() and len(cifras) <= 15


def _normalizar_url(texto):
    """Devuelve (url, nota). Lo que no sea una URL válida pasa a las notas."""
    if not texto:
        return "", ""
    if "://" in texto:
        url = texto
    elif FORMA_DE_HOST.fullmatch(texto):
        url = f"https://{texto}"
    else:
        return "", f"Sitio: {texto}"
    if len(url) > 200:
        return "", f"Sitio: {texto}"
    try:
        URLValidator()(url)
    except ValidationError:
        return "", f"Sitio: {texto}"
    return url, ""


def _url_parece_clave(texto):
    """Valor de la columna URL que no sirvió como dirección y tiene forma de contraseña.

    No cuentan las descripciones (llevan espacios: "Oficina principal"), lo que ya
    es una dirección (esquema, www., IP o \\\\ruta) ni los nombres de equipo sin
    dominio (SRV-DC01), que en una columna "Servidor" son lo normal.
    """
    return not (
        any(c.isspace() for c in texto)
        or "://" in texto
        or PARECE_DIRECCION.match(texto)
        or HOST_SIN_PUNTO.fullmatch(texto)
    )


def _url_dudosa(texto, url):
    """True si la URL se armó anteponiendo https:// a algo como "Clave.Segura".

    Solo dos partes y una terminación que no es de dominio conocido: puede ser una
    contraseña corrida de columna, que quedaría como URL visible para el rol lectura.
    """
    if not url or "://" in texto or PARECE_DIRECCION.match(texto):
        return False
    host = urlsplit(url).hostname or ""
    partes = host.split(".")
    return len(partes) == 2 and len(partes[1]) != 2 and partes[1] not in DOMINIOS_CONOCIDOS


def _fecha(valor):
    """Devuelve (iso | None, error). El error no repite el valor: puede ser una clave."""
    if isinstance(valor, datetime):
        return valor.date().isoformat(), ""
    if isinstance(valor, date):
        return valor.isoformat(), ""
    texto = _texto(valor).strip()
    if not texto:
        return None, ""
    for formato in FORMATOS_FECHA:
        try:
            return datetime.strptime(texto, formato).date().isoformat(), ""
        except ValueError:
            continue
    return None, "Fecha de expiración no válida: usa dd/mm/aaaa o aaaa-mm-dd"


def clave_duplicado(titulo, usuario):
    return (titulo.strip().lower(), usuario.strip().lower())


def _validar(datos, error_clave, estado_crudo, fecha_cruda):
    """Devuelve el motivo del error o "" si la fila es válida."""
    if not datos["titulo"]:
        return "Falta el título"
    if error_clave:
        return error_clave
    if not datos["secreto"].strip():
        return "Falta la contraseña"
    if datos["secreto"] in CLAVES_EJEMPLO:
        return "Fila de ejemplo de la plantilla: bórrala o reemplázala"
    if len(datos["titulo"]) > 150:
        return "El título supera 150 caracteres"
    if len(datos["usuario"]) > 150:
        return "El usuario supera 150 caracteres"
    estado = _sin_tildes(estado_crudo).lower()
    if estado and estado not in ESTADOS:
        # Sin repetir el valor: con las columnas corridas, aquí puede estar la clave.
        return "Estado no válido: usa Activa, Expirada o Revocada (o déjalo vacío)"
    _, error_fecha = _fecha(fecha_cruda)
    return error_fecha


def _revisar_tamano(archivo):
    """Rechaza el .xlsx si se infla al descomprimirlo, antes de que openpyxl lo abra.

    Se cuenta lo que ocupa cada parte del zip ya descomprimida (zipfile no entrega
    más bytes de los declarados): un sharedStrings.xml de pocos KB comprimido puede
    ser cientos de MB de XML que openpyxl carga entero.
    """
    with zipfile.ZipFile(archivo) as paquete:
        tamanos = [parte.file_size for parte in paquete.infolist()]
    archivo.seek(0)
    if max(tamanos, default=0) > MAX_BYTES_PARTE or sum(tamanos) > MAX_BYTES_DESCOMPRIMIDOS:
        raise ArchivoInvalido(
            "El archivo ocupa demasiado al abrirlo (descomprimido pasa de "
            f"{MAX_BYTES_DESCOMPRIMIDOS // 2**20} MB, o de {MAX_BYTES_PARTE // 2**20} MB en una de sus partes). "
            "Copia solo las contraseñas a la plantilla y súbela."
        )


class _SoloPrimeraHoja(ExcelReader):
    """Lo mismo que openpyxl.load_workbook, pero solo prepara la primera hoja de datos.

    Al abrir, openpyxl recorre cada hoja declarada (y lee las de gráficos): un archivo
    de pocos KB con cientos de hojas que apuntan a la misma parte tarda minutos. Aquí
    solo se lee la primera hoja, así que es la única que se prepara.
    """

    def read_worksheets(self):
        primera = next(
            (hoja for hoja, rel in self.parser.find_sheets()
             if rel.target in self.valid_files and "chartsheet" not in rel.Type),
            None,
        )
        self.parser.sheets = [primera] if primera else []
        super().read_worksheets()

    def read(self):
        with warnings.catch_warnings():
            # Los nombres definidos (áreas de impresión…) de las hojas que no se
            # preparan no tienen dónde ir; openpyxl lo avisa y no importa aquí.
            warnings.filterwarnings("ignore", message="Defined names for sheet index", category=UserWarning)
            super().read()


def _abrir_libro(archivo):
    _revisar_tamano(archivo)
    lector = _SoloPrimeraHoja(archivo, read_only=True, data_only=True, keep_links=False)
    lector.read()
    return lector.wb


def _leer_hoja(archivo):
    """Devuelve (hoja, columnas, nombres, ignoradas, filas) de la primera hoja.

    hoja: su nombre; columnas: campo → índice; nombres: campo → encabezado tal como
    viene en el archivo; ignoradas: [(índice, texto)] de encabezados no reconocidos o
    repetidos; filas: [(número de fila en Excel, celdas)] sin las filas vacías.
    """
    wb = _abrir_libro(archivo)
    try:
        ws = wb.worksheets[0]
        # Sin esto openpyxl confía en la dimensión que declara el archivo, que
        # una sola celda con formato lejos de los datos vuelve gigantesca.
        ws.reset_dimensions()
        encabezado = next(ws.iter_rows(min_row=1, max_row=1, max_col=MAX_COLUMNAS, values_only=True), ())
        columnas, nombres, ignoradas = {}, {}, []
        for indice, valor in enumerate(encabezado):
            if valor is None or str(valor).strip() == "":
                continue
            campo = _campo_de(valor)
            if campo and campo not in columnas:
                columnas[campo] = indice
                nombres[campo] = str(valor).strip()
            else:
                ignoradas.append((indice, str(valor).strip()))

        ancho = max(columnas.values(), default=0) + 1
        filas = []
        # openpyxl entrega una tupla por cada número de fila desde la 2, también por
        # los huecos: el número sale de la posición y acota el recorrido.
        for numero, celdas in enumerate(ws.iter_rows(min_row=2, max_col=ancho), start=2):
            if numero > MAX_FILA_EXCEL:
                raise ArchivoInvalido(
                    f"El archivo tiene filas después de la última de Excel ({_miles(MAX_FILA_EXCEL)}). "
                    "¿Es un .xlsx válido?"
                )
            if all(_texto(getattr(c, "value", None)).strip() == "" for c in celdas):
                continue
            filas.append((numero, celdas))
            if len(filas) > MAX_FILAS:
                raise ArchivoInvalido(f"El archivo tiene más de {MAX_FILAS} filas con datos.")
        return ws.title, columnas, nombres, ignoradas, filas
    finally:
        wb.close()


def _miles(numero):
    return f"{numero:,}".replace(",", ".")


def _enumerar(textos):
    textos = list(textos)
    return textos[0] if len(textos) == 1 else ", ".join(textos[:-1]) + " y " + textos[-1]


def _advertencias(hoja, columnas, ignoradas):
    advertencias = []
    faltantes = [nombre for campo, nombre in COLUMNAS_OBLIGATORIAS if campo not in columnas]
    for nombre in faltantes:
        advertencias.append(f'No se encontró la columna "{nombre}" en la hoja «{hoja}» (se lee solo la primera hoja).')
    if not faltantes and "usuario" not in columnas:
        advertencias.append(
            'No se encontró la columna "Usuario": las filas con el mismo título se tratarán como repetidas.'
        )
    if ignoradas and faltantes:
        # Sin las columnas obligatorias, la fila 1 seguramente son datos (quizá una
        # clave): se nombran las columnas por su letra, sin mostrar los valores.
        letras = _enumerar(get_column_letter(indice + 1) for indice, _ in ignoradas)
        cuales = f"Se ignoró la columna {letras}" if len(ignoradas) == 1 else f"Se ignoraron las columnas {letras}"
        advertencias.append(
            f"{cuales}: la fila 1 no tiene encabezados que se reconozcan "
            "(ver la hoja Instrucciones de la plantilla)."
        )
    elif ignoradas:
        nombres = ", ".join(f"«{texto}»" for _, texto in ignoradas)
        advertencias.append(
            f"Se ignoraron columnas que no se reconocen o están repetidas: {nombres}. "
            "Si tienen datos que necesitas, renombra el encabezado (ver la hoja Instrucciones)."
        )
    return advertencias, faltantes


def _avisos_de_columnas_cruzadas(nombres, con_clave, claves_como_direccion, urls_escritas, urls_como_clave):
    """Avisos que bloquean el archivo entero porque las columnas parecen pegadas en otro orden.

    Regla (la clave terminaría en texto plano en URL o en Notas):
    - Contraseñas: al menos MIN_FILAS_ORDEN filas con clave y la mitad o más con
      forma de dirección (http://, https://, www., \\\\ruta o IP). Con menos filas no
      se bloquea: una sola clave "www.Casa2024" es una clave.
    - URL: al menos MIN_FILAS_ORDEN URL escritas y la mitad o más con forma de
      contraseña según _url_parece_clave (no cuentan los nombres de equipo sin
      dominio como SRV-DC01 ni las descripciones con espacios).
    Cada aviso nombra la columna, cuántos valores y por qué.
    """
    avisos = []
    if con_clave >= MIN_FILAS_ORDEN and claves_como_direccion * 2 >= con_clave:
        avisos.append(
            f"La columna «{nombres.get('secreto', 'Contraseña')}»: {claves_como_direccion} de {con_clave} "
            "contraseñas parecen direcciones (empiezan por http://, www. o son una IP). ¿Pegaste las "
            "columnas en otro orden? Corrige el archivo y vuelve a subirlo: no se importó nada."
        )
    if urls_escritas >= MIN_FILAS_ORDEN and urls_como_clave * 2 >= urls_escritas:
        avisos.append(
            f"La columna «{nombres.get('url', 'URL')}» (URL): {urls_como_clave} de {urls_escritas} valores "
            "no son direcciones y parecen contraseñas (sin espacios y con símbolos). ¿Pegaste las columnas "
            "en otro orden? Si esa columna no tiene direcciones, cámbiale el encabezado (por ejemplo, a "
            "«Notas»). No se importó nada."
        )
    return avisos


def leer_credenciales(archivo, vault):
    """Lee la primera hoja del Excel y valida cada fila contra la bóveda.

    Returns:
        (filas, advertencias, hoja). Cada fila es un dict serializable a JSON con
        fila, titulo, usuario, url, url_a_notas, notas, estado,
        fecha_expiracion, secreto_cifrado, resultado ("nueva" | "ya_existe" |
        "repetida" | "error") y detalle. hoja es el nombre de la hoja leída.

    Columnas corridas: si la clave de una fila parece una dirección y su URL parece
    una clave (o se armó con una terminación desconocida, como "Clave.Segura"), esa
    fila sale con error y sin URL ni notas. Si una columna entera parece cruzada
    (ver _avisos_de_columnas_cruzadas) no se importa nada.

    Aviso "Revisa" por fila (no bloquea) cuando el valor de la columna URL no es una
    dirección y queda como nota: la nota no se cifra, así que si esa columna traía
    contraseñas (Contraseña y Servidor intercambiadas, con claves sin símbolos como
    "Casa2024", que tienen la forma de un nombre de equipo) la clave real quedaría en
    claro. No se bloquea porque "SRV-DC01" en una columna "Servidor" es legítimo.
    """
    hoja, columnas, nombres, ignoradas, crudas = _leer_hoja(archivo)
    advertencias, faltantes = _advertencias(hoja, columnas, ignoradas)

    existentes = {
        clave_duplicado(t, u)
        for t, u in Credencial.objects.filter(vault=vault).values_list("titulo", "usuario")
    }
    vistas = set()
    fernet = build_fernet()
    filas = []
    con_clave = claves_como_direccion = urls_escritas = urls_como_clave = 0
    for numero, celdas in crudas:
        def celda(campo, celdas=celdas):
            indice = columnas.get(campo)
            return celdas[indice] if indice is not None and indice < len(celdas) else None

        def valor(campo):
            return getattr(celda(campo), "value", None)

        celda_clave = celda("secreto")
        valor_clave = getattr(celda_clave, "value", None)
        error_clave = ""
        if valor_clave is not None:
            error_clave = _error_de_la_clave(
                valor_clave,
                getattr(celda_clave, "data_type", None),
                getattr(celda_clave, "number_format", "General"),
            )
        datos = {
            "titulo": _texto(valor("titulo")).strip(),
            "usuario": _texto(valor("usuario")).strip(),
            "secreto": "" if error_clave else _texto(valor_clave),
        }
        estado_crudo = _texto(valor("estado")).strip()
        fecha_cruda = valor("fecha_expiracion")
        url_original = _texto(valor("url")).strip()
        url, nota_url = _normalizar_url(url_original)
        notas = "\n".join(t for t in (_texto(valor("notas")).strip(), nota_url) if t)
        fecha_iso, _ = _fecha(fecha_cruda)

        clave_es_direccion = bool(datos["secreto"].strip()) and bool(PARECE_DIRECCION.match(datos["secreto"].strip()))
        url_parece_clave = bool(nota_url) and _url_parece_clave(url_original)
        url_dudosa = _url_dudosa(url_original, url)
        if datos["secreto"].strip():
            con_clave += 1
            claves_como_direccion += clave_es_direccion
        if url_original:
            urls_escritas += 1
            urls_como_clave += url_parece_clave

        registro = {
            "fila": numero,
            "titulo": datos["titulo"],
            "usuario": datos["usuario"],
            "url": url,
            "url_a_notas": bool(nota_url),
            "notas": notas,
            "estado": _sin_tildes(estado_crudo).lower() or "activa",
            "fecha_expiracion": fecha_iso,
            "secreto_cifrado": "",
            "resultado": "nueva",
            "detalle": "",
        }

        if faltantes:
            error = "Falta la columna " + " y ".join(f'"{n}"' for n in faltantes)
        else:
            error = _validar(datos, error_clave, estado_crudo, fecha_cruda)
        if not error and clave_es_direccion and (url_parece_clave or url_dudosa):
            error = "La contraseña parece una dirección y la URL una contraseña: ¿columnas corridas en esta fila?"
            registro.update(url="", url_a_notas=False, notas="")
        clave = clave_duplicado(datos["titulo"], datos["usuario"])
        if error:
            registro.update(resultado="error", detalle=error)
        elif clave in existentes:
            registro.update(resultado="ya_existe", detalle="Ya existe: mismo título y usuario")
        elif clave in vistas:
            registro.update(resultado="repetida", detalle="Repetida en el archivo")
        else:
            vistas.add(clave)
            registro["secreto_cifrado"] = fernet.encrypt(datos["secreto"].encode("utf-8")).decode("utf-8")
            revisar = []
            if isinstance(valor_clave, (int, float)):
                revisar.append("en Excel era un número")
            if url_dudosa:
                revisar.append("la URL no termina en un dominio conocido (.com, .co, .local…), ¿es una contraseña?")
            if nota_url and not any(c.isspace() for c in url_original):
                revisar.append(
                    f"«{nombres.get('url', 'URL')}» no es una dirección y quedó como nota sin cifrar, "
                    "¿es una contraseña? (muévela a la columna de contraseña)"
                )
            if revisar:
                registro["detalle"] = "Revisa: " + "; ".join(revisar)
        filas.append(registro)

    cruzadas = _avisos_de_columnas_cruzadas(nombres, con_clave, claves_como_direccion, urls_escritas, urls_como_clave)
    if cruzadas:
        advertencias.extend(cruzadas)
        for registro in filas:
            registro.update(resultado="error", detalle="Revisa el orden de las columnas",
                            secreto_cifrado="", url="", url_a_notas=False, notas="")

    return filas, advertencias, hoja


def tamano_en_sesion(datos):
    """Bytes que ocupan los datos en la sesión (el mismo JSON que usa Django)."""
    return len(json.dumps(datos, separators=(",", ":")))


def resumen(filas):
    return {
        "nuevas": sum(1 for f in filas if f["resultado"] == "nueva"),
        "ya_existen": sum(1 for f in filas if f["resultado"] in ("ya_existe", "repetida")),
        "errores": sum(1 for f in filas if f["resultado"] == "error"),
        "total": len(filas),
    }


def plantilla_workbook():
    """Plantilla Excel para importar credenciales (columnas en formato Texto)."""
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Plantilla Contraseñas"
    borde = Border(*(Side(style="thin") for _ in range(4)))

    for fila in range(1, MAX_FILAS + 2):
        for columna in range(1, len(ENCABEZADOS_PLANTILLA) + 1):
            ws.cell(row=fila, column=columna).number_format = "@"

    for columna, texto in enumerate(ENCABEZADOS_PLANTILLA, 1):
        celda = ws.cell(row=1, column=columna, value=texto)
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill(start_color="0156A6", end_color="0156A6", fill_type="solid")
        celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        celda.border = borde
        ws.column_dimensions[openpyxl.utils.get_column_letter(columna)].width = 26

    for fila, ejemplo in enumerate(EJEMPLOS_PLANTILLA, 2):
        for columna, valor in enumerate(ejemplo, 1):
            celda = ws.cell(row=fila, column=columna, value=valor or None)
            celda.border = borde

    hoja = wb.create_sheet("Instrucciones")
    for fila, linea in enumerate(INSTRUCCIONES_PLANTILLA, 1):
        celda = hoja.cell(row=fila, column=1, value=linea)
        if fila == 1:
            celda.font = Font(bold=True, size=14, color="0156A6")
    hoja.column_dimensions["A"].width = 110
    return wb
