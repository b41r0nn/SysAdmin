import io
from unittest.mock import patch

from django.contrib.auth import get_user_model
from django.test import Client, TestCase, override_settings
from django.urls import reverse
from openpyxl import load_workbook

from inventario.models import Activo
from mantenimiento.models import OrdenMantenimiento
from notificaciones.models import Notificacion, NotificacionEmail
from soporte.models import Ticket
from soporte.views import _aviso_email_ticket

CustomUser = get_user_model()


def _user(rol):
    return CustomUser.objects.create_user(
        username=f"sopo_{rol}_{CustomUser.objects.count()}",
        password="x",
        rol=rol,
    )


def _activo(serial="SN-SOP-001"):
    return Activo.objects.create(
        serial=serial,
        tipo_dispositivo="escritorio",
        marca="HP",
        modelo="Pro",
        estado="disponible",
    )


def _ticket(solicitante, **kw):
    defaults = {"asunto": "Monitor no enciende", "prioridad": "media"}
    defaults.update(kw)
    return Ticket.objects.create(solicitante=solicitante, **defaults)


# ─── Acceso ────────────────────────────────────────────────────────


class TicketsLoginTests(TestCase):
    def test_lista_requiere_login(self):
        resp = self.client.get(reverse("soporte:lista"))
        self.assertEqual(resp.status_code, 302)

    def test_crear_requiere_login(self):
        resp = self.client.get(reverse("soporte:crear"))
        self.assertEqual(resp.status_code, 302)


# ─── CRUD ──────────────────────────────────────────────────────────


class TicketsCRUDTests(TestCase):
    def setUp(self):
        self.user = _user("superadmin")
        self.client = Client()
        self.client.force_login(self.user)
        self.activo = _activo()

    def test_lista_returns_200(self):
        resp = self.client.get(reverse("soporte:lista"))
        self.assertEqual(resp.status_code, 200)

    def test_crear_ticket_post(self):
        resp = self.client.post(
            reverse("soporte:crear"),
            {
                "asunto": "No abre el sistema",
                "prioridad": "alta",
                "descripcion": "Error al iniciar sesion",
            },
        )
        self.assertEqual(resp.status_code, 302)
        ticket = Ticket.objects.get(asunto="No abre el sistema")
        self.assertEqual(ticket.solicitante, self.user)
        self.assertEqual(ticket.estado, "abierto")

    def test_editar_ticket(self):
        ticket = _ticket(self.user)
        resp = self.client.post(
            reverse("soporte:editar", args=[ticket.pk]),
            {"asunto": "Asunto nuevo", "prioridad": "baja", "descripcion": "Actualizado"},
        )
        self.assertEqual(resp.status_code, 302)
        ticket.refresh_from_db()
        self.assertEqual(ticket.asunto, "Asunto nuevo")
        self.assertEqual(ticket.prioridad, "baja")

    def test_detalle_returns_200(self):
        ticket = _ticket(self.user)
        resp = self.client.get(reverse("soporte:detalle", args=[ticket.pk]))
        self.assertEqual(resp.status_code, 200)

    def test_crear_no_muestra_campos_tecnico(self):
        resp = self.client.get(reverse("soporte:crear"))
        self.assertEqual(resp.status_code, 200)
        self.assertNotContains(resp, "Acciones realizadas")
        self.assertNotContains(resp, "Tiempo empleado")

    def test_editar_muestra_campos_tecnico(self):
        ticket = _ticket(self.user)
        resp = self.client.get(reverse("soporte:editar", args=[ticket.pk]))
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, "Acciones realizadas")
        self.assertContains(resp, "Tiempo empleado (minutos)")


# ─── Permisos por rol ──────────────────────────────────────────────


class TicketsPermisosTests(TestCase):
    def setUp(self):
        self.user = _user("lectura")
        self.client = Client()
        self.client.force_login(self.user)
        self.activo = _activo()

    def test_lectura_puede_ver_lista(self):
        resp = self.client.get(reverse("soporte:lista"))
        self.assertEqual(resp.status_code, 200)

    def test_lectura_puede_crear_ticket(self):
        resp = self.client.get(reverse("soporte:crear"))
        self.assertEqual(resp.status_code, 200)

    def test_lectura_no_puede_asignar(self):
        ticket = _ticket(self.user)
        resp = self.client.post(reverse("soporte:asignar", args=[ticket.pk]))
        self.assertEqual(resp.status_code, 302)

    def test_lectura_no_puede_escalar(self):
        ticket = _ticket(self.user)
        resp = self.client.get(reverse("soporte:escalar", args=[ticket.pk]))
        self.assertEqual(resp.status_code, 302)


# ─── Asignacion y notificacion ─────────────────────────────────────


class AsignacionNotificacionTests(TestCase):
    def setUp(self):
        self.user = _user("superadmin")
        self.client = Client()
        self.client.force_login(self.user)
        self.tecnico = _user("tecnico")
        self.ticket = _ticket(self.user)

    def test_asignar_notifica_al_tecnico(self):
        resp = self.client.post(
            reverse("soporte:asignar", args=[self.ticket.pk]),
            {"asignado_a": self.tecnico.pk},
        )
        self.assertEqual(resp.status_code, 302)
        self.ticket.refresh_from_db()
        self.assertEqual(self.ticket.asignado_a, self.tecnico)
        self.assertEqual(self.ticket.estado, "en_proceso")
        n = Notificacion.objects.filter(
            usuario=self.tecnico,
            tipo="aviso",
            objetokey=f"ticket:{self.ticket.pk}",
        )
        self.assertEqual(n.count(), 1)

    def test_reasignar_no_duplica_notificacion(self):
        self.client.post(
            reverse("soporte:asignar", args=[self.ticket.pk]),
            {"asignado_a": self.tecnico.pk},
        )
        self.client.post(
            reverse("soporte:asignar", args=[self.ticket.pk]),
            {"asignado_a": self.tecnico.pk},
        )
        n = Notificacion.objects.filter(
            usuario=self.tecnico,
            tipo="aviso",
            objetokey=f"ticket:{self.ticket.pk}",
        )
        self.assertEqual(n.count(), 1)


# ─── Escalamiento a orden de mantenimiento ─────────────────────────


class EscalarTicketTests(TestCase):
    def setUp(self):
        self.user = _user("superadmin")
        self.client = Client()
        self.client.force_login(self.user)
        self.activo = _activo()
        self.ticket = _ticket(self.user, prioridad="alta")

    def test_escalar_crea_orden_con_ticket(self):
        resp = self.client.post(
            reverse("soporte:escalar", args=[self.ticket.pk]),
            {
                "activo": self.activo.pk,
                "prioridad": "alta",
                "tecnico_asignado": "Carlos",
            },
        )
        self.assertEqual(resp.status_code, 302)
        orden = OrdenMantenimiento.objects.get(ticket=self.ticket)
        self.assertEqual(orden.activo, self.activo)
        self.assertEqual(orden.tipo, "correctivo")
        self.assertEqual(orden.estado, "abierta")
        self.assertEqual(orden.prioridad, "alta")
        self.assertIn(f"#{self.ticket.pk}", orden.descripcion)
        self.ticket.refresh_from_db()
        self.assertEqual(self.ticket.estado, "escalado")

    def test_escalar_no_duplica_orden(self):
        self.client.post(
            reverse("soporte:escalar", args=[self.ticket.pk]),
            {"activo": self.activo.pk, "prioridad": "media"},
        )
        resp = self.client.post(
            reverse("soporte:escalar", args=[self.ticket.pk]),
            {"activo": self.activo.pk, "prioridad": "media"},
        )
        self.assertEqual(resp.status_code, 302)
        self.assertEqual(OrdenMantenimiento.objects.filter(ticket=self.ticket).count(), 1)


# ─── Avisos por email en cambios de estado ─────────────────────────


class TicketEmailTriggersTests(TestCase):
    def setUp(self):
        self.user = _user("superadmin")
        self.client = Client()
        self.client.force_login(self.user)
        self.tecnico = _user("tecnico")
        self.ticket = _ticket(self.user)

    @patch("soporte.views.encolar_email")
    def test_asignar_encola_email(self, mock):
        resp = self.client.post(
            reverse("soporte:asignar", args=[self.ticket.pk]),
            {"asignado_a": self.tecnico.pk},
        )
        self.assertEqual(resp.status_code, 302)
        mock.assert_called_once()
        args, _ = mock.call_args
        self.assertEqual(args[0], self.ticket.solicitante.email)

    @patch("soporte.views.encolar_email")
    def test_cambiar_estado_encola_email(self, mock):
        resp = self.client.post(
            reverse("soporte:cambiar_estado", args=[self.ticket.pk]),
            {"estado": "en_proceso"},
        )
        self.assertEqual(resp.status_code, 302)
        mock.assert_called_once()
        args, _ = mock.call_args
        self.assertEqual(args[0], self.ticket.solicitante.email)

    @patch("soporte.views.encolar_email")
    def test_cambiar_al_mismo_estado_no_reenvia(self, mock):
        self.ticket.estado = "en_proceso"
        self.ticket.save(update_fields=["estado"])
        resp = self.client.post(
            reverse("soporte:cambiar_estado", args=[self.ticket.pk]),
            {"estado": "en_proceso"},
        )
        self.assertEqual(resp.status_code, 302)
        mock.assert_not_called()

    def test_aviso_email_sin_email_no_rompe(self):
        self.user.email = ""
        self.user.save(update_fields=["email"])
        self.ticket.solicitante = self.user
        self.ticket.save(update_fields=["solicitante"])
        _aviso_email_ticket(self.ticket, "en_proceso")
        self.assertEqual(NotificacionEmail.objects.count(), 0)


# ─── Formulario público ───────────────────────────────────────────


class ReportePublicoTests(TestCase):
    URL = "/soporte/reportar-publico/"

    def setUp(self):
        self.admin = _user("admin")
        self.tecnico = _user("tecnico")

    def test_get_sin_login_200(self):
        resp = self.client.get(self.URL)
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, "Reportar falla")

    def test_post_sin_login_crea_ticket_publico(self):
        resp = self.client.post(self.URL, {
            "nombre": "María Pérez",
            "area": "Recursos Humanos",
            "contacto": "maria@redihos.local",
            "descripcion": "La impresora del 3er piso no imprime.",
        })
        self.assertEqual(resp.status_code, 200)
        ticket = Ticket.objects.get()
        self.assertIsNone(ticket.solicitante)
        self.assertEqual(ticket.nombre_solicitante, "María Pérez")
        self.assertEqual(ticket.area_solicitante, "Recursos Humanos")
        self.assertEqual(ticket.contacto_solicitante, "maria@redihos.local")
        self.assertContains(resp, f"#{ticket.pk}")

    def test_post_guarda_tipo_dispositivo_y_numero_serie(self):
        resp = self.client.post(self.URL, {
            "nombre": "Juan",
            "area": "Ventas",
            "contacto": "",
            "tipo_dispositivo": "portatil",
            "numero_serie_etiqueta": "SN-PORT-999",
            "descripcion": "Portátil no enciende.",
        })
        self.assertEqual(resp.status_code, 200)
        ticket = Ticket.objects.get()
        self.assertEqual(ticket.tipo_dispositivo, "portatil")
        self.assertEqual(ticket.numero_serie_etiqueta, "SN-PORT-999")

    def test_post_tipo_dispositivo_opcional(self):
        resp = self.client.post(self.URL, {
            "nombre": "Luis",
            "area": "Sistemas",
            "contacto": "",
            "tipo_dispositivo": "",
            "numero_serie_etiqueta": "",
            "descripcion": "No enciende.",
        })
        self.assertEqual(resp.status_code, 200)
        ticket = Ticket.objects.get()
        self.assertEqual(ticket.tipo_dispositivo, "")
        self.assertEqual(ticket.numero_serie_etiqueta, "")

    def test_post_honeypot_lleno_no_crea_ticket(self):
        resp = self.client.post(self.URL, {
            "nombre": "Bot",
            "area": "Spam",
            "contacto": "",
            "descripcion": "spam",
            "sitio_web": "http://spam.example.com",
        })
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(Ticket.objects.count(), 0)

    def test_aviso_ti_solo_bandeja_sin_email(self):
        with patch("soporte.views.encolar_email") as mock_email:
            self.client.post(self.URL, {
                "nombre": "Carlos",
                "area": "Compras",
                "contacto": "",
                "descripcion": "PC no arranca.",
            })
        mock_email.assert_not_called()
        for destinatario in (self.admin, self.tecnico):
            self.assertTrue(Notificacion.objects.filter(
                usuario=destinatario,
                tipo="aviso",
                objetokey__startswith="ticket:",
            ).exists())

    def test_aviso_no_suena_para_quien_no_es_ti(self):
        operario = _user("operario")
        self.client.post(self.URL, {
            "nombre": "Carlos",
            "area": "Compras",
            "contacto": "",
            "descripcion": "PC no arranca.",
        })
        self.assertFalse(Notificacion.objects.filter(usuario=operario).exists())

    @override_settings(SOPORTE_REPORTE_PUBLICO_COOLDOWN_SEGUNDOS=2)
    def test_cooldown_bloquea_ticket_segundo_intento(self):
        payload = {
            "nombre": "Ana",
            "area": "Logística",
            "contacto": "",
            "descripcion": "Primer reporte.",
        }
        self.client.post(self.URL, payload)
        self.assertEqual(Ticket.objects.count(), 1)

        # Mismo client, misma sesión: debe bloquear.
        resp = self.client.post(self.URL, {
            "nombre": "Ana",
            "area": "Logística",
            "contacto": "",
            "descripcion": "Segundo reporte inmediato.",
        })
        self.assertEqual(Ticket.objects.count(), 1)
        self.assertContains(resp, "Esperá un momento")

    @override_settings(SOPORTE_REPORTE_PUBLICO_COOLDOWN_SEGUNDOS=0)
    def test_cooldown_cero_permite_multiples_reportes(self):
        for i in range(3):
            self.client.post(self.URL, {
                "nombre": f"Usuario {i}",
                "area": "Sistemas",
                "contacto": "",
                "descripcion": f"Reporte {i}.",
            })
        self.assertEqual(Ticket.objects.count(), 3)


# ─── Navegación ─────────────────────────────────────────────────


class SidebarTests(TestCase):
    def test_reportar_falla_no_aparece_en_nav_autenticado(self):
        user = _user("admin")
        self.client.force_login(user)
        resp = self.client.get(reverse("soporte:lista"))
        self.assertEqual(resp.status_code, 200)
        self.assertNotContains(resp, "mantenimiento:reportar")


class LoginLinkTests(TestCase):
    def test_login_muestra_link_publico(self):
        resp = self.client.get(reverse("accounts:login"))
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, "¿Necesitás reportar una falla?")
        self.assertContains(resp, reverse("soporte:reportar_publico"))

    def test_link_publico_accesible_sin_cuenta(self):
        resp = self.client.get(reverse("soporte:reportar_publico"))
        self.assertEqual(resp.status_code, 200)


# ─── Edición técnica (acciones + tiempo) ─────────────────────────


class EdicionTecnicaTests(TestCase):
    def setUp(self):
        self.tecnico = _user("tecnico")
        self.client.force_login(self.tecnico)

    def _ticket_en_proceso(self):
        ticket = _ticket(self.tecnico)
        ticket.estado = "en_proceso"
        ticket.save(update_fields=["estado"])
        return ticket

    def test_tecnico_puede_editar_acciones_y_tiempo(self):
        ticket = self._ticket_en_proceso()
        resp = self.client.post(
            reverse("soporte:editar", args=[ticket.pk]),
            {
                "asunto": ticket.asunto,
                "descripcion": ticket.descripcion,
                "prioridad": ticket.prioridad,
                "acciones_realizadas": "Se actualizó el BIOS y quedó funcionando.",
                "tiempo_empleado_minutos": "45",
            },
        )
        self.assertEqual(resp.status_code, 302)
        ticket.refresh_from_db()
        self.assertEqual(ticket.acciones_realizadas, "Se actualizó el BIOS y quedó funcionando.")
        self.assertEqual(ticket.tiempo_empleado_minutos, 45)

    def test_editar_sin_tiempo_lo_deja_null(self):
        ticket = self._ticket_en_proceso()
        resp = self.client.post(
            reverse("soporte:editar", args=[ticket.pk]),
            {
                "asunto": ticket.asunto,
                "descripcion": ticket.descripcion,
                "prioridad": ticket.prioridad,
                "acciones_realizadas": "",
                "tiempo_empleado_minutos": "",
            },
        )
        self.assertEqual(resp.status_code, 302)
        ticket.refresh_from_db()
        self.assertIsNone(ticket.tiempo_empleado_minutos)


# ─── Export Excel ─────────────────────────────────────────────────


class TicketExcelTests(TestCase):
    def setUp(self):
        self.user = _user("admin")
        self.tecnico = _user("tecnico")
        self.client.force_login(self.user)

    def test_export_sin_session_requiere_login(self):
        c = Client()
        resp = c.get(reverse("soporte:excel"))
        self.assertEqual(resp.status_code, 302)

    def test_export_responde_xlsx(self):
        _ticket(self.user)
        resp = self.client.get(reverse("soporte:excel"))
        self.assertEqual(resp.status_code, 200)
        self.assertIn(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            resp["Content-Type"],
        )
        self.assertIn("tickets_soporte.xlsx", resp["Content-Disposition"])

    def test_export_ticket_publico_no_rompe(self):
        Ticket.objects.create(
            asunto="Falla reportada: pc lento",
            descripcion="pc lento",
            nombre_solicitante="Ana",
            area_solicitante="Contabilidad",
            contacto_solicitante="",
        )
        resp = self.client.get(reverse("soporte:excel"))
        self.assertEqual(resp.status_code, 200)

    def _valor_metrica(self, resp, metrica):
        wb = load_workbook(io.BytesIO(resp.content))
        ws = wb["Resumen"]
        for row in ws.iter_rows():
            if row[0].value == metrica:
                return row[1].value
        return None

    def test_export_tiempo_real_cuando_existe(self):
        t1 = _ticket(self.user)
        t1.estado = "resuelto"
        t1.tiempo_empleado_minutos = 30
        t1.save(update_fields=["estado", "tiempo_empleado_minutos"])
        t2 = _ticket(self.user, asunto="Segundo ticket")
        t2.estado = "cerrado"
        t2.tiempo_empleado_minutos = 60
        t2.save(update_fields=["estado", "tiempo_empleado_minutos"])
        t3 = _ticket(self.user, asunto="Tercer ticket")
        t3.estado = "resuelto"
        t3.tiempo_empleado_minutos = None
        t3.save(update_fields=["estado", "tiempo_empleado_minutos"])

        resp = self.client.get(reverse("soporte:excel"))
        valor = self._valor_metrica(resp, "Tiempo promedio de resolución")
        # Promedio de los valores reales cargados (30+60)/2 = 45 min; el tercero
        # sin cargar NO se usa para el promedio real.
        self.assertEqual(valor, "45 min")

    def test_export_tiempo_fallback_proxy_sin_real(self):
        t1 = _ticket(self.user)
        t1.estado = "resuelto"
        t1.tiempo_empleado_minutos = None
        t1.save(update_fields=["estado", "tiempo_empleado_minutos"])

        resp = self.client.get(reverse("soporte:excel"))
        valor = self._valor_metrica(resp, "Tiempo promedio de resolución")
        self.assertIn("días", valor)

    def test_export_tiempo_real_incluye_cero_minutos(self):
        t1 = _ticket(self.user)
        t1.estado = "resuelto"
        t1.tiempo_empleado_minutos = 0
        t1.save(update_fields=["estado", "tiempo_empleado_minutos"])
        t2 = _ticket(self.user, asunto="Segundo ticket")
        t2.estado = "cerrado"
        t2.tiempo_empleado_minutos = 60
        t2.save(update_fields=["estado", "tiempo_empleado_minutos"])

        resp = self.client.get(reverse("soporte:excel"))
        valor = self._valor_metrica(resp, "Tiempo promedio de resolución")
        # (0+60)/2 = 30 min: el ticket resuelto en 0 minutos cuenta en el
        # promedio real y NO cae al proxy.
        self.assertEqual(valor, "30 min")