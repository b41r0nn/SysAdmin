from datetime import timedelta

from django.conf import settings
from django.contrib import messages
from django.contrib.auth import get_user_model
from django.contrib.auth.decorators import login_required
from django.core.paginator import Paginator
from django.db.models import Q
from django.http import HttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.urls import reverse
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from accounts.permisos import requiere_permiso
from mantenimiento.models import OrdenMantenimiento
from notificaciones.services import aviso_usuario, encolar_email
from reports.views import (
    _adjust_column_widths,
    _apply_redihos_data_style,
    _apply_redihos_header,
    AZUL,
    AZUL_CLARO,
    AZUL_OSCURO,
    BLANCO,
    GRIS_TXT,
)

from .forms import EscalarTicketForm, TicketForm, TicketPublicoForm
from .models import ESTADOS_TICKET, PRIORIDADES_TICKET, Ticket


def _notificar_ti_ticket(ticket):
    """Aviso en bandeja (sin email) a roles de TI al crear un ticket público."""
    titulo = f"Ticket público #{ticket.pk} · {ticket.nombre_solicitante}"
    tipo_incluye = ""
    if ticket.tipo_dispositivo:
        tipo_incluye = f"Equipo: {ticket.get_tipo_dispositivo_display()}. "
    mensaje = (
        f"{ticket.nombre_solicitante} reportó: \"{ticket.asunto}\". "
        f"Área: {ticket.area_solicitante or '—'}. {tipo_incluye}"
        f"Contacto: {ticket.contacto_solicitante or '—'}"
    )
    link = reverse("soporte:detalle", args=[ticket.pk])
    ti = get_user_model().objects.filter(
        is_active=True, rol__in={"superadmin", "admin", "tecnico"}
    )
    for user in ti:
        aviso_usuario(user, titulo, mensaje, link=link, objetokey=f"ticket:{ticket.pk}")


def _cooldown_reporte_publico(request):
    """Devuelve True si el usuario debe esperar antes de enviar otro ticket.

    El límite se guarda en sesión; si no hay sesión, se permite.
    En desarrollo/tests el cooldown es 0 por defecto.
    """
    segundos = getattr(settings, "SOPORTE_REPORTE_PUBLICO_COOLDOWN_SEGUNDOS", 0)
    if not segundos:
        return False
    if not request.session.session_key:
        return False
    ultimo = request.session.get("ultimo_ticket_publico")
    if not ultimo:
        return False
    try:
        ultimo_dt = timezone.datetime.fromisoformat(ultimo)
    except (TypeError, ValueError):
        return False
    return (timezone.now() - ultimo_dt).total_seconds() < segundos


def reportar_publico(request):
    """Formulario público (LAN interna, sin login) para reportar una falla."""
    if request.method == "POST":
        # Honeypot: si el campo oculto viene lleno, descartar silenciosamente
        # respondiendo éxito, sin crear ticket ni revelar la detección.
        if request.POST.get("sitio_web"):
            return render(request, "soporte/reporte_publico_exito.html", {
                "ticket": None,
            })

        if _cooldown_reporte_publico(request):
            form = TicketPublicoForm(request.POST)
            messages.error(
                request,
                "Esperá un momento antes de enviar otro reporte. "
                "Si es urgente, contactá directamente a Sistemas."
            )
            return render(request, "soporte/reporte_publico.html", {"form": form})

        form = TicketPublicoForm(request.POST)
        if form.is_valid():
            ticket = Ticket.objects.create(
                asunto=f"Falla reportada: {form.cleaned_data['descripcion'][:100]}",
                descripcion=form.cleaned_data["descripcion"],
                solicitante=None,
                nombre_solicitante=form.cleaned_data["nombre"],
                area_solicitante=form.cleaned_data["area"],
                contacto_solicitante=form.cleaned_data["contacto"],
                tipo_dispositivo=form.cleaned_data["tipo_dispositivo"],
                numero_serie_etiqueta=form.cleaned_data["numero_serie_etiqueta"],
            )
            _notificar_ti_ticket(ticket)
            request.session["ultimo_ticket_publico"] = timezone.now().isoformat()
            return render(request, "soporte/reporte_publico_exito.html", {
                "ticket": ticket,
            })
    else:
        form = TicketPublicoForm()

    return render(request, "soporte/reporte_publico.html", {"form": form})


def _aviso_email_ticket(ticket, estado_nuevo):
    """Encola email al solicitante cuando cambia el estado de un ticket."""
    if not ticket.solicitante:
        return
    encolar_email(
        ticket.solicitante.email,
        f"Ticket #{ticket.pk}: {dict(ESTADOS_TICKET)[estado_nuevo]}",
        f"Tu ticket \"{ticket.asunto}\" cambió a estado "
        f"\"{dict(ESTADOS_TICKET)[estado_nuevo]}\".\nSoporte TI",
    )


@requiere_permiso("soporte", "lectura")
def lista_tickets(request):
    qs = Ticket.objects.select_related("solicitante", "asignado_a")

    estado = request.GET.get("estado", "").strip()
    prioridad = request.GET.get("prioridad", "").strip()
    q = request.GET.get("q", "").strip()

    if estado:
        qs = qs.filter(estado=estado)
    if prioridad:
        qs = qs.filter(prioridad=prioridad)
    if q:
        qs = qs.filter(
            Q(asunto__icontains=q) | Q(solicitante__username__icontains=q)
        )

    stats = {
        "total": qs.count(),
        "abiertos": qs.filter(estado="abierto").count(),
        "en_proceso": qs.filter(estado="en_proceso").count(),
        "resueltos": qs.filter(estado__in=["resuelto", "cerrado", "escalado"]).count(),
    }

    page_number = request.GET.get("page", "1")
    paginator = Paginator(qs.order_by("-fecha_creacion"), 20)
    page_obj = paginator.get_page(page_number)

    context = {
        "tickets": page_obj.object_list,
        "page_obj": page_obj,
        "filtro_estado": estado,
        "filtro_prioridad": prioridad,
        "filtro_q": q,
        "stats": stats,
        "estados": ESTADOS_TICKET,
        "prioridades": PRIORIDADES_TICKET,
    }
    return render(request, "soporte/lista_tickets.html", context)


@login_required
def crear_ticket(request):
    if request.method == "POST":
        form = TicketForm(request.POST)
        if form.is_valid():
            ticket = form.save(commit=False)
            ticket.solicitante = request.user
            ticket.save()
            messages.success(request, "Ticket creado correctamente.")
            return redirect("soporte:detalle", pk=ticket.pk)
        messages.error(request, "Corrige los errores del formulario.")
    else:
        form = TicketForm()

    return render(request, "soporte/form_ticket.html", {
        "form": form,
        "titulo": "Nuevo ticket",
    })


@requiere_permiso("soporte", "lectura")
def detalle_ticket(request, pk):
    ticket = get_object_or_404(
        Ticket.objects.select_related("solicitante", "asignado_a"), pk=pk
    )
    asignables = get_user_model().objects.filter(is_active=True).order_by("username")
    return render(request, "soporte/detalle_ticket.html", {
        "ticket": ticket,
        "asignables": asignables,
    })


@requiere_permiso("soporte", "escritura")
def editar_ticket(request, pk):
    ticket = get_object_or_404(Ticket, pk=pk)
    if request.method == "POST":
        form = TicketForm(request.POST, instance=ticket)
        if form.is_valid():
            ticket = form.save()
            messages.success(request, "Ticket actualizado.")
            return redirect("soporte:detalle", pk=ticket.pk)
        messages.error(request, "Corrige los errores del formulario.")
    else:
        form = TicketForm(instance=ticket)

    return render(request, "soporte/form_ticket.html", {
        "form": form,
        "titulo": f"Editar ticket #{ticket.pk}",
        "ticket": ticket,
    })


@requiere_permiso("soporte", "escritura")
def asignar_ticket(request, pk):
    ticket = get_object_or_404(Ticket, pk=pk)
    if request.method == "POST":
        user_id = request.POST.get("asignado_a", "").strip()
        user = None
        if user_id:
            user = get_object_or_404(get_user_model(), pk=user_id)
        estado_anterior = ticket.estado
        ticket.asignado_a = user
        if estado_anterior == "abierto":
            ticket.estado = "en_proceso"
        ticket.save(update_fields=["asignado_a", "estado", "fecha_actualizacion"])
        if ticket.estado != estado_anterior:
            _aviso_email_ticket(ticket, ticket.estado)
        if user:
            aviso_usuario(
                user,
                f"Ticket asignado · #{ticket.pk}",
                f"Te asignaron el ticket \"{ticket.asunto}\".",
                link=reverse("soporte:detalle", args=[ticket.pk]),
                objetokey=f"ticket:{ticket.pk}",
            )
            messages.success(request, f"Ticket asignado a {user.username}.")
        else:
            messages.success(request, "El ticket quedó sin asignación.")
    return redirect("soporte:detalle", pk=ticket.pk)


@requiere_permiso("soporte", "escritura")
def cambiar_estado(request, pk):
    ticket = get_object_or_404(Ticket, pk=pk)
    if request.method == "POST":
        nuevo = request.POST.get("estado", "")
        if nuevo in dict(ESTADOS_TICKET):
            anterior = ticket.estado
            ticket.estado = nuevo
            ticket.save(update_fields=["estado", "fecha_actualizacion"])
            if nuevo != anterior:
                _aviso_email_ticket(ticket, nuevo)
            messages.success(request, f"Ticket marcado como {dict(ESTADOS_TICKET)[nuevo]}.")
    return redirect("soporte:detalle", pk=ticket.pk)


@requiere_permiso("soporte", "escritura")
def escalar_ticket(request, pk):
    ticket = get_object_or_404(Ticket, pk=pk)
    escalado = ticket.ordenes.first()
    if escalado:
        messages.info(request, "Este ticket ya fue escalado.")
        return redirect("mantenimiento:detalle_orden", pk=escalado.pk)

    if request.method == "POST":
        form = EscalarTicketForm(request.POST, initial={"prioridad": ticket.prioridad})
        if form.is_valid():
            orden = OrdenMantenimiento.objects.create(
                ticket=ticket,
                activo=form.cleaned_data["activo"],
                tipo="correctivo",
                estado="abierta",
                prioridad=form.cleaned_data["prioridad"],
                tecnico_asignado=form.cleaned_data["tecnico_asignado"],
                descripcion=f"Escalado desde ticket #{ticket.pk}: {ticket.asunto}",
            )
            ticket.estado = "escalado"
            ticket.save(update_fields=["estado", "fecha_actualizacion"])
            _aviso_email_ticket(ticket, "escalado")
            messages.success(request, "Ticket escalado a orden de mantenimiento.")
            return redirect("mantenimiento:detalle_orden", pk=orden.pk)
        messages.error(request, "Corrige los errores del formulario.")
    else:
        form = EscalarTicketForm(initial={"prioridad": ticket.prioridad})

    return render(request, "soporte/form_escalar.html", {
        "form": form,
        "ticket": ticket,
    })


def _nombre_solicitante(ticket):
    if ticket.nombre_solicitante:
        return ticket.nombre_solicitante
    if ticket.solicitante:
        return ticket.solicitante.username
    return "—"


@requiere_permiso("soporte", "lectura")
def tickets_excel(request):
    """Exporta estadísticas y detalle de tickets a Excel (patrón reports)."""
    qs = Ticket.objects.select_related("solicitante", "asignado_a")

    estado = request.GET.get("estado", "").strip()
    prioridad = request.GET.get("prioridad", "").strip()
    if estado:
        qs = qs.filter(estado=estado)
    if prioridad:
        qs = qs.filter(prioridad=prioridad)

    wb = Workbook()
    ws = wb.active
    ws.title = "Tickets"

    # ── Resumen de estadísticas ─────────────────────────────────────────────
    base_qs = qs
    total = base_qs.count()
    por_estado = {e: base_qs.filter(estado=e).count() for e, _ in ESTADOS_TICKET}
    por_prioridad = {p: base_qs.filter(prioridad=p).count() for p, _ in PRIORIDADES_TICKET}

    cerrados = base_qs.filter(estado__in=["resuelto", "cerrado"])
    tiempo_promedio = None
    uso_real = False
    if cerrados.count():
        con_tiempo = [t.tiempo_empleado_minutos for t in cerrados if t.tiempo_empleado_minutos is not None]
        if con_tiempo:
            # Tiempo real cargado por el técnico: promedio en minutos.
            tiempo_promedio = sum(con_tiempo) / len(con_tiempo)
            uso_real = True
        else:
            # Fallback: proxy fecha_actualizacion − fecha_creacion.
            diffs = [
                (t.fecha_actualizacion - t.fecha_creacion).total_seconds()
                for t in cerrados
            ]
            seg_prom = sum(diffs) / len(diffs)
            tiempo_promedio = timedelta(seconds=seg_prom)

    resumen = [
        ("Total de tickets", total),
        ("Abiertos", por_estado["abierto"]),
        ("En proceso", por_estado["en_proceso"]),
        ("Resueltos", por_estado["resuelto"] + por_estado["cerrado"] + por_estado["escalado"]),
        ("Cerrados", por_estado["cerrado"]),
        ("Escalados", por_estado["escalado"]),
        ("Prioridad alta", por_prioridad["alta"]),
        ("Prioridad media", por_prioridad["media"]),
        ("Prioridad baja", por_prioridad["baja"]),
    ]
    if tiempo_promedio is not None:
        if uso_real:
            resumen.append(("Tiempo promedio de resolución", f"{tiempo_promedio:.0f} min"))
        else:
            horas = tiempo_promedio.total_seconds() / 3600
            dias = horas / 24
            resumen.append(("Tiempo promedio de resolución", f"{dias:.1f} días ({horas:.1f} h)"))
    else:
        resumen.append(("Tiempo promedio de resolución", "Sin tickets cerrados"))

    # Resumen en una hoja aparte (2 columnas: encabezado manual, porque la
    # helper corporativa reserva columnas A/B para el logo y exige >= 3 columnas).
    ws_resumen = wb.create_sheet("Resumen")
    ws_resumen.title = "Resumen"
    thin_side = Side(style="thin", color="D1D5DB")
    border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

    ws_resumen.merge_cells("A1:B1")
    ws_resumen["A1"] = "REDIHOS S.A.S."
    ws_resumen["A1"].fill = PatternFill(start_color=AZUL_OSCURO, end_color=AZUL_OSCURO, fill_type="solid")
    ws_resumen["A1"].font = Font(name="Calibri", bold=True, size=16, color=BLANCO)
    ws_resumen["A1"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws_resumen.row_dimensions[1].height = 48

    ws_resumen.merge_cells("A2:B2")
    ws_resumen["A2"] = f"Estadísticas de Tickets — {timezone.now().strftime('%d/%m/%Y')}"
    ws_resumen["A2"].fill = PatternFill(start_color=AZUL_CLARO, end_color=AZUL_CLARO, fill_type="solid")
    ws_resumen["A2"].font = Font(name="Calibri", bold=True, size=12, color=AZUL_OSCURO)

    resumen_headers = ["Métrica", "Valor"]
    header_row = 4
    for col_idx, h in enumerate(resumen_headers, start=1):
        cell = ws_resumen.cell(row=header_row, column=col_idx, value=h)
        cell.fill = PatternFill(start_color=AZUL, end_color=AZUL, fill_type="solid")
        cell.font = Font(name="Calibri", bold=True, size=10, color=BLANCO)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border

    data_start = header_row + 1
    for row_idx, (metrica, valor) in enumerate(resumen, start=data_start):
        ws_resumen.cell(row=row_idx, column=1, value=metrica)
        ws_resumen.cell(row=row_idx, column=2, value=valor)
    ws_resumen.column_dimensions["A"].width = 36
    ws_resumen.column_dimensions["B"].width = 30
    for row_idx in range(data_start, data_start + len(resumen)):
        for col_idx in (1, 2):
            cell = ws_resumen.cell(row=row_idx, column=col_idx)
            cell.font = Font(name="Calibri", size=10, color=GRIS_TXT)
            cell.border = border
            cell.alignment = Alignment(
                horizontal="left" if col_idx == 1 else "center",
                vertical="center",
                indent=1,
            )

    # ── Detalle de tickets ──────────────────────────────────────────────────
    headers = [
        "Nº", "Asunto", "Solicitante", "Área", "Contacto", "Tipo", "Nº serie/etiqueta",
        "Prioridad", "Estado", "Asignado a", "Creado",
    ]
    start_row = _apply_redihos_header(ws, headers, "Tickets de Soporte")
    for idx, t in enumerate(qs.order_by("-fecha_creacion"), start=start_row):
        row = [
            t.pk,
            t.asunto,
            _nombre_solicitante(t),
            t.area_solicitante,
            t.contacto_solicitante,
            t.get_tipo_dispositivo_display() if t.tipo_dispositivo else "—",
            t.numero_serie_etiqueta or "—",
            t.get_prioridad_display(),
            t.get_estado_display(),
            t.asignado_a.username if t.asignado_a else "—",
            t.fecha_creacion.strftime("%d/%m/%Y %H:%M"),
        ]
        ws.append(row)

    _apply_redihos_data_style(ws, start_row, len(headers))
    _adjust_column_widths(ws, headers)

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = 'attachment; filename="tickets_soporte.xlsx"'
    wb.save(response)
    return response