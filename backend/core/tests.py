import os
import tempfile
from datetime import timedelta
from io import StringIO
from pathlib import Path
from types import SimpleNamespace

from django.apps import apps
from django.conf import settings
from django.contrib.admin.models import LogEntry
from django.contrib.auth import get_user_model
from django.contrib.contenttypes.models import ContentType
from django.core.management import call_command
from django.test import Client, TestCase, override_settings
from django.urls import reverse
from django.utils import timezone

from accounts.permisos import PERMISOS_POR_ROL
from core.utils import get_client_ip
from core.views import MEDIA_PERMISOS

CustomUser = get_user_model()


def _user(rol, username=None):
    return CustomUser.objects.create_user(
        username=username or f"core_{rol}_{CustomUser.objects.count()}",
        password="x",
        rol=rol,
    )


class ClientIPTests(TestCase):
    def test_toma_ultimo_valor_de_forwarded_for(self):
        request = SimpleNamespace(META={"HTTP_X_FORWARDED_FOR": "9.9.9.9, 10.0.0.5"})
        self.assertEqual(get_client_ip(request), "10.0.0.5")

    def test_sin_forwarded_usa_remote_addr(self):
        request = SimpleNamespace(META={"REMOTE_ADDR": "192.168.1.10"})
        self.assertEqual(get_client_ip(request), "192.168.1.10")

    def test_forwarded_vacio_cae_a_remote_addr(self):
        request = SimpleNamespace(
            META={"HTTP_X_FORWARDED_FOR": "", "REMOTE_ADDR": "192.168.1.10"}
        )
        self.assertEqual(get_client_ip(request), "192.168.1.10")

    def test_request_none(self):
        self.assertIsNone(get_client_ip(None))

    def test_request_sin_meta(self):
        self.assertIsNone(get_client_ip(SimpleNamespace()))


class VaultKeyCheckTests(TestCase):
    def _run(self):
        from core.checks import check_passwords_encryption_key

        return check_passwords_encryption_key(None)

    @override_settings(PASSWORDS_ENCRYPTION_KEY="")
    def test_w001_si_falta_clave(self):
        self.assertEqual([e.id for e in self._run()], ["sysadmin.W001"])

    @override_settings(PASSWORDS_ENCRYPTION_KEY="clave-fernet-de-prueba")
    def test_sin_warning_con_clave(self):
        self.assertEqual(self._run(), [])


class DashboardTests(TestCase):
    def setUp(self):
        self.user = _user("superadmin")
        self.client = Client()
        self.client.force_login(self.user)

    def test_login_required(self):
        self.client.logout()
        resp = self.client.get(reverse("core:dashboard"))
        self.assertEqual(resp.status_code, 302)
        self.assertIn(reverse("accounts:login"), resp.url)

    def test_status_200(self):
        resp = self.client.get(reverse("core:dashboard"))
        self.assertEqual(resp.status_code, 200)


def _cerrar_archivo(resp):
    """Libera el handle del archivo de un FileResponse antes de borrar tmpdir.

    Consumir streaming_content deja que el wrapper del test client cierre el
    archivo con close_old_connections desconectado. resp.close() directo
    dispara request_finished y, con PostgreSQL, cierra la conexión de BD en
    mitad del TestCase.
    """
    b"".join(resp.streaming_content)


class MediaProtegidoTests(TestCase):
    def setUp(self):
        self.tmpdir = tempfile.TemporaryDirectory()
        self.media_root = Path(self.tmpdir.name)

        # Crear archivos de prueba dentro de cada prefijo permitido
        (self.media_root / "usuarios" / "fotos").mkdir(parents=True)
        (self.media_root / "inventario" / "fotos").mkdir(parents=True)
        (self.media_root / "inventario" / "actas" / "pdf").mkdir(parents=True)
        (self.media_root / "mantenimiento" / "fotos").mkdir(parents=True)
        (self.media_root / "documentos" / "2026" / "10").mkdir(parents=True)

        (self.media_root / "usuarios" / "fotos" / "x.png").write_bytes(
            b"\x89PNG\r\n\x1a\n"  # magic PNG
        )
        (self.media_root / "inventario" / "fotos" / "a.png").write_bytes(
            b"\x89PNG\r\n\x1a\n"
        )
        (self.media_root / "documentos" / "2026" / "10" / "malicioso.html").write_text(
            "<html></html>"
        )
        (self.media_root / "otro").mkdir(parents=True, exist_ok=True)
        (self.media_root / "otro" / "a.png").write_bytes(b"\x89PNG\r\n\x1a\n")

        self.user_tecnico = _user("tecnico", username="media_tecnico")
        self.user_lectura = _user("lectura", username="media_lectura")
        self.user_admin = _user("admin", username="media_admin")

    def tearDown(self):
        self.tmpdir.cleanup()

    def _url(self, ruta):
        return reverse("media", kwargs={"ruta": ruta})

    def _override_media(self, accel=False):
        kw = {"MEDIA_ROOT": str(self.media_root)}
        if accel:
            kw["MEDIA_ACCEL_REDIRECT"] = True
        return override_settings(**kw)

    def test_anonimo_redirige_a_login(self):
        with self._override_media():
            resp = Client().get(self._url("usuarios/fotos/x.png"))
        self.assertEqual(resp.status_code, 302)
        self.assertIn(reverse("accounts:login"), resp.url)

    def test_usuario_con_permiso_ve_imagen(self):
        client = Client()
        client.force_login(self.user_tecnico)
        with self._override_media():
            resp = client.get(self._url("inventario/fotos/a.png"))
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp["Content-Type"], "image/png")
        self.assertTrue(resp["Content-Disposition"].startswith("inline"))
        _cerrar_archivo(resp)

    def test_usuario_sin_permiso_obtiene_404(self):
        # Todos los roles reales tienen al menos lectura en los prefijos registrados.
        # Forzamos un rol con permiso "ninguno" sobre inventario para este test.
        original = PERMISOS_POR_ROL["tecnico"]["inventario"]
        PERMISOS_POR_ROL["tecnico"]["inventario"] = "ninguno"
        try:
            client = Client()
            client.force_login(self.user_tecnico)
            with self._override_media():
                resp = client.get(self._url("inventario/fotos/a.png"))
            self.assertEqual(resp.status_code, 404)
        finally:
            PERMISOS_POR_ROL["tecnico"]["inventario"] = original

    def test_path_traversal_devuelve_404(self):
        client = Client()
        client.force_login(self.user_tecnico)
        rutas = [
            "usuarios/fotos/../../.env",
            "usuarios/fotos/..%2f..%2f.env",
            "//etc/passwd",
            "usuarios/fotos/../otro/a.png",
        ]
        for ruta in rutas:
            with self.subTest(ruta=ruta), self._override_media():
                resp = client.get(self._url(ruta))
            self.assertEqual(resp.status_code, 404, ruta)

    def test_prefijo_no_registrado_devuelve_404(self):
        client = Client()
        client.force_login(self.user_tecnico)
        with self._override_media():
            resp = client.get(self._url("otro/a.png"))
        self.assertEqual(resp.status_code, 404)

    def test_html_se_fuerza_a_descarga(self):
        client = Client()
        client.force_login(self.user_admin)
        with self._override_media():
            resp = client.get(self._url("documentos/2026/10/malicioso.html"))
        self.assertEqual(resp.status_code, 200)
        self.assertTrue(resp["Content-Disposition"].startswith("attachment"))
        _cerrar_archivo(resp)

    def test_x_accel_redirect_en_produccion(self):
        client = Client()
        client.force_login(self.user_tecnico)
        with self._override_media(accel=True):
            resp = client.get(self._url("inventario/fotos/a.png"))
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp["X-Accel-Redirect"], "/_protegido/inventario/fotos/a.png")
        self.assertEqual(resp.content, b"")

    def test_superadmin_accede_a_cualquier_prefijo(self):
        superadmin = _user("superadmin", username="media_super")
        client = Client()
        client.force_login(superadmin)
        with self._override_media():
            resp = client.get(self._url("documentos/2026/10/malicioso.html"))
        self.assertEqual(resp.status_code, 200)
        _cerrar_archivo(resp)

    def test_todos_los_upload_to_estan_registrados(self):
        """Si se añade un nuevo FileField/ImageField sin registrar su prefijo,
        este test falla y obliga a actualizar MEDIA_PERMISOS."""
        registrados = tuple(pref for pref, _ in MEDIA_PERMISOS)

        for model in apps.get_models():
            for field in model._meta.get_fields():
                if field.get_internal_type() not in ("FileField", "ImageField"):
                    continue
                upload_to = getattr(field, "upload_to", "")
                if not isinstance(upload_to, str):
                    continue
                # Quitar placeholders de fecha (documentos/%Y/%m/ -> documentos/)
                prefijo = upload_to.split("%")[0].rstrip("/") + "/"
                self.assertTrue(
                    any(prefijo.startswith(r) for r in registrados),
                    f"{model.__name__}.{field.name} -> upload_to='{upload_to}' no está "
                    f"cubierto por MEDIA_PERMISOS (prefijo calculado: '{prefijo}')",
                )


class VendorAssetsTests(TestCase):
    """Los assets front (bootstrap, htmx, fullcalendar, fonts) se sirven
    localmente desde static/vendor; ninguna plantilla debe apuntar a un CDN."""

    HOSTS_CDN = (
        "cdn.jsdelivr.net",
        "unpkg.com",
        "googleapis.com",
        "gstatic.com",
        "cdnjs",
    )

    def test_plantillas_sin_cdn_externos(self):
        base = Path(settings.BASE_DIR)
        directorios = [base / "templates"]
        directorios += [d / "templates" for d in base.iterdir() if (d / "templates").is_dir()]
        encontrados = []
        for directorio in directorios:
            for plantilla in directorio.rglob("*.html"):
                contenido = plantilla.read_text(encoding="utf-8")
                for host in self.HOSTS_CDN:
                    if host in contenido:
                        encontrados.append(f"{plantilla.relative_to(base)} -> {host}")
        self.assertEqual(encontrados, [], f"Plantillas con CDN externo: {encontrados}")


class LimpiarHistorialCommandTests(TestCase):
    def test_dry_run_no_borra_nada(self):
        from administracion.models import RegistroAuditoria

        registro = RegistroAuditoria.objects.create(modulo="inventario", accion="crear")
        RegistroAuditoria.objects.all().update(fecha=timezone.now() - timedelta(days=400))

        out = StringIO()
        call_command("limpiar_historial", dias=365, dry_run=True, stdout=out)

        self.assertTrue(RegistroAuditoria.objects.filter(pk=registro.pk).exists())
        salida = out.getvalue()
        self.assertIn("simulación", salida)
        self.assertIn("Total registros afectados", salida)

    def test_borra_solo_lo_antiguo(self):
        from administracion.models import RegistroAuditoria
        from notificaciones.models import Notificacion, NotificacionEmail
        from yule.models import SincronizacionLog

        usuario = _user("superadmin")
        viejo = timezone.now() - timedelta(days=400)
        nuevo = timezone.now()

        audit_viejo = RegistroAuditoria.objects.create(modulo="inventario", accion="crear")
        audit_nuevo = RegistroAuditoria.objects.create(modulo="inventario", accion="crear")
        RegistroAuditoria.objects.filter(pk=audit_viejo.pk).update(fecha=viejo)

        sync = SincronizacionLog.objects.create()
        SincronizacionLog.objects.filter(pk=sync.pk).update(fecha_inicio=viejo)

        content_type = ContentType.objects.get_for_model(RegistroAuditoria)
        log_viejo = LogEntry.objects.create(
            user=usuario,
            content_type=content_type,
            object_id="1",
            object_repr="x",
            action_flag=1,
            action_time=viejo,
        )

        notif_vieja = Notificacion.objects.create(usuario=usuario, tipo="aviso", titulo="t", leida=True)
        Notificacion.objects.filter(pk=notif_vieja.pk).update(fecha_creacion=viejo)
        notif_nueva = Notificacion.objects.create(usuario=usuario, tipo="aviso", titulo="t", leida=True)
        Notificacion.objects.filter(pk=notif_nueva.pk).update(fecha_creacion=nuevo)

        mail_viejo = NotificacionEmail.objects.create(
            destinatario="a@b.co", asunto="s", cuerpo="c", enviado=True, fecha_enviado=viejo
        )
        mail_nuevo = NotificacionEmail.objects.create(
            destinatario="a@b.co", asunto="s2", cuerpo="c", enviado=True, fecha_enviado=nuevo
        )

        call_command("limpiar_historial", dias=365, stdout=StringIO())

        self.assertFalse(RegistroAuditoria.objects.filter(pk=audit_viejo.pk).exists())
        self.assertTrue(RegistroAuditoria.objects.filter(pk=audit_nuevo.pk).exists())
        self.assertFalse(SincronizacionLog.objects.filter(pk=sync.pk).exists())
        self.assertFalse(LogEntry.objects.filter(pk=log_viejo.pk).exists())
        self.assertFalse(Notificacion.objects.filter(pk=notif_vieja.pk).exists())
        self.assertTrue(Notificacion.objects.filter(pk=notif_nueva.pk).exists())
        self.assertFalse(NotificacionEmail.objects.filter(pk=mail_viejo.pk).exists())
        self.assertTrue(NotificacionEmail.objects.filter(pk=mail_nuevo.pk).exists())

    def test_borra_media_huerfana_y_respeta_referenciada(self):
        from inventario.tests import crear_activo

        with tempfile.TemporaryDirectory() as tmp:
            ref_dir = Path(tmp) / "inventario" / "fotos"
            ref_dir.mkdir(parents=True)
            referenciada = ref_dir / "keep.jpg"
            referenciada.write_bytes(b"ref")
            huerfana = Path(tmp) / "orphan.jpg"
            huerfana.write_bytes(b"huerfana")

            antiguo = (timezone.now() - timedelta(days=60)).timestamp()
            os.utime(referenciada, (antiguo, antiguo))
            os.utime(huerfana, (antiguo, antiguo))

            crear_activo("SN-MEDIA-1", foto_activo="inventario/fotos/keep.jpg")

            with override_settings(MEDIA_ROOT=tmp):
                call_command("limpiar_historial", dias=365, media_dias=30, stdout=StringIO())

            self.assertTrue(referenciada.exists())
            self.assertFalse(huerfana.exists())


class GuardarExcelTests(TestCase):
    def test_texto_con_igual_no_queda_como_formula(self):
        from io import BytesIO

        from openpyxl import Workbook, load_workbook

        from core.utils import guardar_excel

        wb = Workbook()
        wb.active.append(["=HYPERLINK(\"http://x\",\"clic\")", "normal"])
        buffer = BytesIO()
        guardar_excel(wb, buffer)
        buffer.seek(0)
        celda = load_workbook(buffer).active["A1"]
        self.assertEqual(celda.data_type, "s")
        self.assertEqual(celda.value, "=HYPERLINK(\"http://x\",\"clic\")")


class SmokeVistasTests(TestCase):
    """Abre con superadmin las páginas de solo lectura de cada módulo.

    Existe porque un import roto dentro de una vista (no en la cabecera del
    módulo) solo falla al abrir esa página: la lista de inventario estuvo rota
    en producción sin que ningún test la tocara.
    """

    RUTAS = [
        "core:dashboard",
        "usuarios:lista", "usuarios:crear", "usuarios:importar", "usuarios:importar_plantilla",
        "inventario:lista", "inventario:software_global", "inventario:crear", "inventario:etiquetas",
        "inventario:catalogo_lista", "inventario:catalogo_crear", "inventario:catalogo_json",
        "inventario:exportar_opciones", "inventario:importar", "inventario:importar_plantilla",
        "reports:index", "reports:activos_por_usuario", "reports:inventario_excel",
        "reports:movimientos_excel", "reports:usuarios_excel", "reports:costos_excel",
        "reports:mantenimiento_excel",
        "mantenimiento:lista_ordenes", "mantenimiento:lista_planes", "mantenimiento:calendario",
        "mantenimiento:calendario_eventos", "mantenimiento:reportar", "mantenimiento:crear_plan",
        "mantenimiento:crear_orden", "mantenimiento:hoja_de_vida_buscar",
        "passwords:index", "passwords:vault_crear",
        "yule:index", "yule:configuracion", "yule:equipos_lista", "yule:equipos_sin_match",
        "yule:historial",
        "documentos:lista", "documentos:crear", "documentos:categorias_lista",
        "administracion:lista_auditoria", "administracion:cuentas", "administracion:crear_cuenta",
        "administracion:configuracion",
        "notificaciones:lista", "notificaciones:cantidad",
        "soporte:lista", "soporte:reportar_publico", "soporte:excel", "soporte:crear",
        "licencias:lista", "licencias:crear", "licencias:exportar_excel",
        "prestamos:lista", "prestamos:crear",
    ]

    def test_paginas_de_lectura_responden_200(self):
        from inventario.models import Activo

        user = get_user_model().objects.create_user(
            username="smoke_admin", password="x", rol="superadmin"
        )
        self.client.force_login(user)
        activo = Activo.objects.create(
            serial="SN-SMOKE", tipo_dispositivo="escritorio", marca="HP",
            modelo="X", estado="disponible",
        )
        rutas = [(r, []) for r in self.RUTAS] + [
            ("inventario:detalle", [activo.pk]),
            ("inventario:software", [activo.pk]),
            ("inventario:editar", [activo.pk]),
            ("inventario:asignar", [activo.pk]),
            ("inventario:qr_publico", [activo.pk]),
            ("mantenimiento:hoja_de_vida", [activo.pk]),
        ]
        for nombre, args in rutas:
            with self.subTest(ruta=nombre):
                resp = self.client.get(reverse(nombre, args=args))
                self.assertEqual(resp.status_code, 200, nombre)
