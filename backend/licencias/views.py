import os
from datetime import date

from django.conf import settings
from django.contrib import messages
from django.db.models import Q
from django.http import HttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.template.loader import render_to_string
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from accounts.permisos import requiere_permiso

from .forms import LicenciaForm
from .models import ESTADOS_LICENCIA, LicenciaSoftware, TIPOS_LICENCIA

LOGO_PATH = os.path.join(settings.BASE_DIR, "static", "img", "logo_redihos_mark.png")


# Mapeo de columnas para el reporte de licencias.
# Cada entrada: key: (label, extractor(licencia))
# El orden define el orden de columnas en el reporte completo.
CAMPOS_LICENCIAS = {
    "id": ("ID", lambda l: l.id),
    "nombre": ("Software", lambda l: l.nombre),
    "version": ("Versión", lambda l: l.version),
    "proveedor": ("Proveedor / editor", lambda l: l.proveedor),
    "clave": ("Clave / serial", lambda l: l.clave),
    "tipo": ("Tipo de licencia", lambda l: l.get_tipo_display()),
    "cantidad": ("Cantidad", lambda l: l.cantidad),
    "estado": ("Estado", lambda l: dict(ESTADOS_LICENCIA)[l.estado_efectivo]),
    "fecha_compra": ("Fecha de compra", lambda l: l.fecha_compra.strftime("%Y-%m-%d") if l.fecha_compra else ""),
    "fecha_vencimiento": ("Fecha de vencimiento", lambda l: l.fecha_vencimiento.strftime("%Y-%m-%d") if l.fecha_vencimiento else ""),
    "costo": ("Costo", lambda l: float(l.costo) if l.costo else None),
    "responsable": ("Responsable", lambda l: l.responsable),
    "cubiertos": ("Equipos cubiertos", lambda l: l.activos.count()),
    "observaciones": ("Observaciones", lambda l: l.observaciones),
    "fecha_creacion": ("Fecha de creación", lambda l: l.fecha_creacion.strftime("%Y-%m-%d")),
}

BADGES_ESTADO = {
    "activa": "success",
    "por_vencer": "warning",
    "vencida": "danger",
    "cancelada": "secondary",
}


def _licencias_filtradas(request):
    qs = LicenciaSoftware.objects.prefetch_related("activos")
    q = request.GET.get("q", "").strip()
    tipo = request.GET.get("tipo", "").strip()
    estado = request.GET.get("estado", "").strip()

    if q:
        qs = qs.filter(Q(nombre__icontains=q) | Q(proveedor__icontains=q) | Q(responsable__icontains=q))
    if tipo:
        qs = qs.filter(tipo=tipo)

    if estado:
        # El estado efectivo se calcula en Python según vencimiento
        objetos = [l for l in qs if l.estado_efectivo == estado]
    else:
        objetos = list(qs)

    return qs, objetos, {"q": q, "tipo": tipo, "estado": estado}


@requiere_permiso("licencias", "lectura")
def lista(request):
    _qs, objetos, filtros = _licencias_filtradas(request)

    stats = {
        "total": len(objetos),
        "activas": sum(1 for l in objetos if l.estado_efectivo == "activa"),
        "por_vencer": sum(1 for l in objetos if l.estado_efectivo == "por_vencer"),
        "vencidas": sum(1 for l in objetos if l.estado_efectivo == "vencida"),
    }

    return render(request, "licencias/lista.html", {
        "licencias": objetos,
        "filtro_q": filtros["q"],
        "filtro_tipo": filtros["tipo"],
        "filtro_estado": filtros["estado"],
        "stats": stats,
        "tipos": TIPOS_LICENCIA,
        "estados": ESTADOS_LICENCIA,
        "badges_estado": BADGES_ESTADO,
    })


@requiere_permiso("licencias", "lectura")
def detalle(request, pk):
    licencia = get_object_or_404(
        LicenciaSoftware.objects.prefetch_related("activos"), pk=pk
    )
    return render(request, "licencias/detalle.html", {
        "licencia": licencia,
        "badge_estado": BADGES_ESTADO[licencia.estado_efectivo],
    })


@requiere_permiso("licencias", "escritura")
def crear(request):
    if request.method == "POST":
        form = LicenciaForm(request.POST)
        if form.is_valid():
            form.save()
            messages.success(request, "Licencia creada correctamente.")
            return redirect("licencias:lista")
        messages.error(request, "Corrige los errores del formulario.")
    else:
        form = LicenciaForm()

    return render(request, "licencias/form.html", {
        "form": form,
        "titulo": "Nueva licencia",
    })


@requiere_permiso("licencias", "escritura")
def editar(request, pk):
    licencia = get_object_or_404(LicenciaSoftware, pk=pk)
    if request.method == "POST":
        form = LicenciaForm(request.POST, instance=licencia)
        if form.is_valid():
            form.save()
            messages.success(request, "Licencia actualizada.")
            return redirect("licencias:detalle", pk=licencia.pk)
        messages.error(request, "Corrige los errores del formulario.")
    else:
        form = LicenciaForm(instance=licencia)

    return render(request, "licencias/form.html", {
        "form": form,
        "titulo": f"Editar — {licencia.nombre}",
        "licencia": licencia,
    })


@requiere_permiso("licencias", "escritura")
def eliminar(request, pk):
    licencia = get_object_or_404(LicenciaSoftware, pk=pk)
    if request.method == "POST":
        nombre = str(licencia)
        licencia.delete()
        messages.warning(request, f"Licencia «{nombre}» eliminada.")
        return redirect("licencias:lista")

    return render(request, "licencias/confirm_delete.html", {
        "titulo": "Eliminar licencia",
        "mensaje": f"Eliminar la licencia «{licencia}»? Esta acción no se puede deshacer.",
        "volver_url": "licencias:detalle",
        "volver_kwargs": {"pk": licencia.pk},
    })


@requiere_permiso("licencias", "lectura")
def exportar_excel(request):
    _qs, objetos, filtros = _licencias_filtradas(request)
    campos = list(CAMPOS_LICENCIAS.keys())

    wb = Workbook()
    ws = wb.active
    ws.title = "Licencias"

    headers = [CAMPOS_LICENCIAS[c][0] for c in campos]
    ws.append(headers)

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="0156A6", end_color="0156A6", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border

    for licencia in objetos:
        fila = []
        for c in campos:
            raw = CAMPOS_LICENCIAS[c][1](licencia)
            if raw is None:
                fila.append("")
            elif c == "costo" and raw:
                fila.append(f"${raw:,.0f} COP")
            else:
                fila.append(raw)
        ws.append(fila)

    from openpyxl.utils import get_column_letter
    for col_idx in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = 26

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = "attachment; filename=licencias.xlsx"
    wb.save(response)
    return response


@requiere_permiso("licencias", "lectura")
def exportar_pdf(request):
    import weasyprint

    _qs, objetos, filtros = _licencias_filtradas(request)
    campos = list(CAMPOS_LICENCIAS.keys())

    filas = []
    for licencia in objetos:
        fila = []
        for c in campos:
            raw = CAMPOS_LICENCIAS[c][1](licencia)
            if c == "estado":
                valor = f'<span class="badge {BADGES_ESTADO[licencia.estado_efectivo]}">{dict(ESTADOS_LICENCIA)[licencia.estado_efectivo]}</span>'
                vacio = False
            elif c == "costo" and raw is not None:
                valor = f"${raw:,.0f} COP"
                vacio = False
            else:
                vacio = raw in (None, "")
                valor = "—" if vacio else raw
            fila.append({"valor": valor, "vacio": vacio})
        filas.append(fila)

    totales = {
        "total": len(objetos),
        "activas": sum(1 for l in objetos if l.estado_efectivo == "activa"),
        "por_vencer": sum(1 for l in objetos if l.estado_efectivo == "por_vencer"),
        "vencidas": sum(1 for l in objetos if l.estado_efectivo == "vencida"),
        "valor": None,
    }
    con_costo = [float(l.costo) for l in objetos if l.costo]
    if con_costo:
        totales["valor"] = f"${sum(con_costo):,.0f} COP"

    html_string = render_to_string("licencias/export_pdf.html", {
        "columnas": [{"label": CAMPOS_LICENCIAS[c][0]} for c in campos],
        "filas": filas,
        "totales": totales,
        "logo_path": LOGO_PATH,
        "generado": timezone.now(),
        "filtros": filtros,
    })
    pdf_bytes = weasyprint.HTML(
        string=html_string, base_url=request.build_absolute_uri("/")
    ).write_pdf()

    response = HttpResponse(pdf_bytes, content_type="application/pdf")
    response["Content-Disposition"] = 'inline; filename="reporte_licencias.pdf"'
    return response